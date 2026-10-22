"""A fictional ten-tab P&O review workbook (invented people and numbers) for the tests."""

from datetime import date

from openpyxl import Workbook


def make_po_workbook(path):
    """A fictional P&O review workbook with the same tab layouts as the real one."""
    import random
    from datetime import timedelta

    rng = random.Random(7)
    stores = ["HM PK LAH Fortress", "HM PK LAH Packages Mall", "HM PK KCH Lucky One"]
    sections = [
        ("01:CGD-FMCG", "CGD - Grocery", "01-CGD"), ("02:Market", "FFD - Bakery", "S054 - Bakery/Pastry"),
        ("02:Market", "FFD - Fruits & Vegetables", "S056 - Fruits & Vegetables"), ("03:Light Household", "LHH - Light Household (Common)", "03-LHH"),
        ("05:Textile", "TXT - Textile (Common)", "05-TXT"),
    ]
    people = []
    n = 0
    for store in stores:
        for dept, section, _ in sections:
            n += 1
            people.append((f"DEMO{n:04d}", f"Example Leader {n:02d}", store, dept, section, "Team Leader" if n % 2 else "Section Manager"))
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(title, rows):
        ws = workbook.create_sheet(title)
        for row in rows:
            ws.append(row)

    rows = []
    for store in stores:
        rows.append(["Store Name", "Dept/Section", "Store Productivity", "Store Productivity LY", "Target Productivity", "Forecasted Productivity "])
        for label, base in [("00 - Overall Store", 100)] + [(code, rng.choice([12, 30, 50, 400])) for _, _, code in sections]:
            rows.append([f"65{stores.index(store)} {store[6:]}", label, round(base * rng.uniform(0.8, 1.2), 1), round(base * rng.uniform(0.8, 1.2), 1), base, round(base * rng.uniform(0.85, 1.15), 1)])
        rows.append([])
    sheet("YTD Productivity (fictional)", rows)

    days = [date(2026, 10, 1) + timedelta(days=i) for i in range(31)]
    codes = ["MOR", "MID", "Evening", "DAY OFF", "MOR", "Evening", "AL"]
    rows = [["Instructions:"], ["", "Morning", "till 11 Consider all in Morning"], [],
            ["Business Unit", "Department ", "Section ", "Employee Number", "Employee Name", "Designation", *days]]
    for emp, name, store, dept, section, role in people:
        rows.append([store, dept, section, emp, name, role, *[codes[(i + int(emp[-2:])) % 7] for i in range(31)]])
    sheet("Roster", rows)

    rows = [["User Status", "Form Template Name", "Subject Employee ID", "Subject Full Name", "Subject Division/Entity", "Subject Position Title",
             "Subject Department", "Subject Pay Grade", "Subject Gender", "Manager Full Name", "HR Business Email", "Overall Performance Rating",
             "Overall Performance Rating Description", "Unadjusted Calculated Overall Performance Rating", "Section Rating", "Section Rating Label"]]
    labels = {2: "Partially Meets Expectations", 3: "Meets expectations", 4: "Exceeds expectations"}
    for emp, name, store, dept, section, role in people:
        r = rng.choice([2, 3, 3, 4])
        rows.append(["Active User", "2025 Performance Journey", emp, name, f"{store} (MAFR-000)", role, f"{section} (MAFR-000.000)",
                     "000-10**-Full Time-Local", "M", "Example Manager", "", r, labels[r], r - 0.2, r, labels[r]])
    sheet("Performance Rating", rows)

    rows = []
    for store in stores:
        rows += [["Report Name: 200-10-05-Country Periodic Store Performance Report"],
                 [f"Store Name: {store}", "", "", "", "", "Report Period : 01/09/2026 - 30/09/2026"], [],
                 ["Section Code Name", "Net Sales"], ["", "Actual", "Forecast", "Budget", "Gth %"]]
        total = budget_total = 0
        for code, name in [("S014", "Grocery [Dry Food]"), ("S054", "Bakery/Pastry"), ("S056", "Fruits & Vegetables"), ("S031", "House-ware"), ("S044", "Men")]:
            actual = rng.randint(5, 90) * 1_000_000
            budget = int(actual * rng.uniform(0.85, 1.15))
            total += actual
            budget_total += budget
            rows.append([f"{code} - {name}", actual, 0, budget, f"{rng.uniform(-10, 25):.1f}%", "0.0%", "0.0%", "", "", "", "", "", "", "", "", "", "", "",
                         f"{rng.uniform(-2, 20):.1f}%", "0.5%", "1.0%", 1000, f"{rng.uniform(5, 25):.1f}%"])
        rows.append(["Total Store", total, 0, budget_total, "12.0%"])
    sheet("Sales", rows)

    rows = [["Details", "Section", "Product Influence", "NPS "]]
    for year in (2025, 2026):
        for sec, cat in [("Bakery ", "Bakery"), ("Market - F&V", "Fruit and Vegetables"), ("TXT", "Menswear")]:
            rows.append([f"Lahore NPS {year}", sec, cat, rng.randint(-20, 90)])
    sheet("NPS", rows)

    header1, header2, header3 = ["Sections"], [""], [""]
    for store in stores:
        header1 += [store, "", "", "", ""]
        header2 += ["Shrinkges", "", "", "Stock Days", ""]
        header3 += ["Budget %", "Actual %", "Act-Bud", "Budget", "Actual"]
    rows = [header1, header2, header3]
    for _, section, _ in sections:
        row = [section]
        for _store in stores:
            b, a = -round(rng.uniform(0.1, 1.5), 2), -round(rng.uniform(0.0, 3.0), 2)
            row += [f"{b}%", f"{a}%", f"{round(a - b, 2)}%", rng.randint(10, 60), rng.randint(10, 80)]
        rows.append(row)
    sheet("Shrinkage & Stock Days", rows)

    rows = [["Sr", "Employee Number", "Business Unit", "Department ", "Section ", "Employee Name", "Designation", "on this role", "Remarks"]]
    for i, (emp, name, store, dept, section, role) in enumerate(people, 1):
        rows.append([i, emp, store, dept, section, name, role, date(2026 - i % 4, 1 + i % 12, 1), f"Leading {section.split(' - ')[-1]} since joining"])
    sheet("In this role", rows)

    sheet("DA", [["Row Labels", "Count of Section"], ["FFD - Bakery", 2], ["HM PK LAH Fortress", 2], ["CGD - Grocery", 1], ["Lucky One", 1], ["Grand Total", 3]])

    rows = [["Emp. No.", "Emp. Name", "Requested Date", "Business Unit", "Dept.", "Sect.", "End of Service Type", "Reason", "Salary Period"]]
    for i in range(18):
        d = date(2026, 9, 28) - timedelta(days=21 * i)
        rows.append([f"DEMOX{i:03d}", f"Example Leaver {i:02d}", d.strftime("%d/%m/%Y"), rng.choice(stores), "Market", rng.choice(["Bakery", "Fruits/Veg", "Butchery"]),
                     rng.choice(["Resignation", "Resignation", "Termination"]), rng.choice(["Found Better Job", "Personal Reasons", "Low Performance"]), d.strftime("%Y%m")])
    sheet("Turnover", rows)

    rows = [["Employee Number", "Business Unit", "Department", "Section", "Employee Name", "Designation ", "Grade", "Gender", "Reporting Manager ",
             "Age ", "Birth Date", "Joining Date", "Total Service Years", "Marital Status", "Local Email", "Local Mobile Number"]]
    for i, (emp, name, store, dept, section, role) in enumerate(people):
        rows.append([emp, store, dept, section, name, role, "10", "Male", people[0][1] if i else "Example Store Manager", "", "", f"01/0{1 + i % 9}/20{15 + i % 9}", "", "", "", ""])
    sheet("Basic Sheet", rows)
    for ws in workbook.worksheets:
        ws.insert_rows(1)
        ws["A1"] = "FICTIONAL TRAINING DATA — NOT COMPANY DATA"
    workbook.save(path)
    return path
