"""Every report pack in every format opens and contains the right things."""

import zipfile

import openpyxl
import pytest
from pptx import Presentation

from pno import engine, exports


def _kw(api, pack):
    m = "2026-09"
    if pack == "scorecard":
        return {"emp": api.dispatch("people", {"month": m, "role": "SM"})["rows"][0]["emp"]}
    if pack == "store":
        return {"store_id": api.dispatch("stores", {"month": m})["rows"][0]["store_id"]}
    if pack == "dept":
        return {"scope": {"dept": "02"}}
    return {}


@pytest.mark.parametrize("pack", list(exports.PACKS))
@pytest.mark.parametrize("fmt", ["pdf", "pptx", "xlsx"])
def test_every_pack_every_format(api, tmp_path, pack, fmt):
    path = tmp_path / f"{pack}.{fmt}"
    r = api.dispatch("export", {"pack": pack, "fmt": fmt, "month": "2026-09", "save_as": str(path), "open_after": False,
                                **_kw(api, pack)})
    assert r.get("ok"), r
    data = path.read_bytes()
    assert len(data) > 2000
    if fmt == "pdf":
        assert data[:5] == b"%PDF-" and b"%%EOF" in data[-1024:]
    elif fmt == "pptx":
        prs = Presentation(str(path))
        assert len(prs.slides) >= 1
    else:
        wb = openpyxl.load_workbook(path)
        assert wb.sheetnames


def test_director_pack_content(api, tmp_path):
    p = tmp_path / "d.pptx"
    api.dispatch("export", {"pack": "director", "fmt": "pptx", "month": "2026-09", "save_as": str(p), "open_after": False})
    prs = Presentation(str(p))
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "September 2026" in text or "Sept 2026" in text
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert charts, "the director deck should carry native, editable charts"
    x = tmp_path / "d.xlsx"
    api.dispatch("export", {"pack": "director", "fmt": "xlsx", "month": "2026-09", "save_as": str(x), "open_after": False})
    with zipfile.ZipFile(x) as z:
        assert any(n.startswith("xl/charts/") for n in z.namelist())


def test_scope_and_hidden_names_reach_the_export(api, db, tmp_path):
    x = tmp_path / "people.xlsx"
    api.dispatch("export", {"pack": "people", "fmt": "xlsx", "month": "2026-09", "scope": {"city": "KCH"}, "hide": True,
                            "save_as": str(x), "open_after": False})
    wb = openpyxl.load_workbook(x)
    values = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value]
    mm = engine.model(db, "2026-09")
    karachi = [p.name for p in mm.people.values() if p.city == "KCH"]
    lahore_only = [p.name for p in mm.people.values() if p.city == "LAH" and p.name not in karachi]
    assert not any(n in values for n in karachi[:30])
    assert not any(n in values for n in lahore_only[:30])
    assert any(v.startswith("Employee ") for v in values)


def test_file_names_are_safe(api, db):
    rep = exports.build_pack(db, "director", "2026-09", {"city": "LAH"})
    name = exports.file_name(rep, "pdf")
    assert name.endswith(".pdf") and not any(ch in name for ch in '\\/:*?"<>|')


def test_dept_pack_needs_a_department(api, tmp_path):
    r = api.dispatch("export", {"pack": "dept", "fmt": "pdf", "month": "2026-09", "save_as": str(tmp_path / "x.pdf"),
                                "open_after": False})
    assert "department" in r["error"].lower()
