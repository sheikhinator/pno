"""Learning Studio: courses sized to the hours, rich lessons, quizzes, AI lessons, assignments, exports and certificates."""

from __future__ import annotations

import json
import zipfile

import pytest

from pno.api import Api
from pno.db import Database
from pno.learning import Learning, clean_lesson, export as LX, library


@pytest.fixture()
def api(tmp_path, monkeypatch):
    monkeypatch.setenv("PNO_HOME", str(tmp_path / "home"))
    return Api(Database(tmp_path / "pno.sqlite3"))


def test_library_is_complete_and_well_formed():
    assert len(library.TOPICS) >= 12
    kinds = set()
    for t in library.TOPICS.values():
        assert t["lessons"] and t["quiz"] and t["activity"]
        for l in t["lessons"]:
            assert l["title"] and l["objectives"] and l["blocks"]
            kinds |= {b["type"] for b in l["blocks"]}
        for q in t["quiz"]:
            if q["type"] in ("mcq", "tf"):
                assert 0 <= q["answer"] < len(q["options"])
            assert q["explain"]
    assert {"para", "points", "formula", "example", "chart", "kpis", "steps", "table", "dodont", "callout"} <= kinds
    for level, keys in library.ROLE_TOPICS.items():
        assert all(k in library.TOPICS for k in keys)


def test_course_fits_the_hours(api):
    for hours, lec in ((2, 20), (6, 25), (10, 30)):
        c = api.dispatch("learn_build", {"level": 3, "hours": hours, "lecture_min": lec})["content"]
        minutes = sum(l["minutes"] for m in c["modules"] for l in m["lessons"]) + 10 * len(c["modules"])
        assert hours * 60 * 0.7 <= minutes <= hours * 60 + lec, (hours, minutes)
        assert all(l["minutes"] == lec for m in c["modules"] for l in m["lessons"])
        assert c["exam"] and all(q["type"] in ("mcq", "tf", "calc") for q in c["exam"])


def test_focus_text_and_topics_pick_the_content(api):
    c = api.dispatch("learn_build", {"level": 1, "hours": 2, "focus": "new cashiers handling refunds and angry customers"})["content"]
    assert {m["key"] for m in c["modules"]} <= {"cash", "customer"}
    c = api.dispatch("learn_build", {"level": 2, "topics": ["shrink"], "hours": 3})["content"]
    assert [m["key"] for m in c["modules"]] == ["shrink"] and c["title"].startswith("Shrinkage")
    c = api.dispatch("learn_build", {"level": 2, "hours": 2, "focus": "forklift licence renewal"})["content"]
    assert c["notes"] and "not in the built-in library" in c["notes"][0]


def test_ai_lessons_are_validated(api, monkeypatch):
    plan = json.dumps({"title": "Ramadan readiness", "summary": "Plan the season", "lessons": ["Forecasting Ramadan demand"]})
    lesson = json.dumps({"title": "Forecasting Ramadan demand", "objectives": ["Plan stock"], "blocks": [
        {"type": "para", "text": "Demand for dates rises sharply."}, {"type": "chart", "kind": "bar", "title": "Dates sales", "labels": ["W1", "W2"],
                                                                      "series": [{"name": "Sales", "values": [10, 30]}]},
        {"type": "chart", "kind": "bar", "title": "broken", "labels": ["a"], "series": [{"name": "x", "values": [1, 2]}]},
        {"type": "video", "url": "x"}], "quiz": [{"type": "mcq", "q": "Q?", "options": ["a", "b"], "answer": 1, "explain": "b"},
                                                  {"type": "mcq", "q": "bad", "options": ["a"], "answer": 5}]})
    answers = iter([(plan, "groq"), (lesson, "groq")])
    monkeypatch.setattr(api.agent, "complete", lambda *a, **k: next(answers))
    c = api.dispatch("learn_build", {"level": 3, "hours": 2, "focus": "Ramadan season readiness"})["content"]
    ai = c["modules"][0]
    assert ai["title"] == "Ramadan readiness" and ai["lessons"][0]["source"] == "AI"
    assert [b["type"] for b in ai["lessons"][0]["blocks"]] == ["para", "chart"]      # broken chart and unknown block dropped
    assert len(ai["quiz"]) == 1
    assert clean_lesson({"title": "x", "blocks": []}) is None


def test_assign_score_certificate_and_tracking(api, tmp_path):
    c = api.dispatch("learn_build", {"level": 2, "hours": 2})
    c = api.dispatch("learn_assign", {"course_id": c["id"], "people": [{"person": "123", "name": "Ayesha Khan", "store": "LAH Fortress"}], "due": "2026-01-01"})
    a = c["assignments"][0]
    api.dispatch("learn_result", {"assignment_id": a["id"], "status": "completed", "score": 86})
    st = api.dispatch("learn_list", {})["stats"]
    assert st["assigned"] == 1 and st["completed"] == 1 and st["avg_score"] == 86 and st["passed"] == 1
    out = tmp_path / "cert.pdf"
    r = api.dispatch("learn_certificate", {"assignment_id": a["id"], "save_as": str(out), "open_after": False})
    assert r["ok"] and out.read_bytes()[:4] == b"%PDF"
    p = api.dispatch("learn_path_save", {"name": "Team Leader path", "level": 2, "course_ids": [c["id"]]})
    assert p["paths"][0]["course_ids"] == [c["id"]]


def test_every_export_format(api, tmp_path):
    c = api.dispatch("learn_build", {"level": 3, "hours": 4})
    for fmt in LX.EXPORTERS:
        out = tmp_path / f"c.{LX.EXT[fmt]}"
        r = api.dispatch("learn_export", {"course_id": c["id"], "fmt": fmt, "save_as": str(out), "open_after": False})
        assert r["ok"] and out.stat().st_size > 4000, fmt
    names = zipfile.ZipFile(tmp_path / "c.zip").namelist()
    assert len(names) == 5 and any(n.endswith(".pptx") for n in names) and any(n.endswith("Answer key.pdf") for n in names)
    from pptx import Presentation
    prs = Presentation(str(tmp_path / "c.pptx"))
    assert len(prs.slides) > 20 and any(s.has_notes_slide and s.notes_slide.notes_text_frame.text for s in prs.slides)
    assert any(sh.has_chart for s in prs.slides for sh in s.shapes)
    import docx
    assert len(docx.Document(str(tmp_path / "c.docx")).paragraphs) > 100


def test_rewrite_and_edit(api, monkeypatch):
    c = api.dispatch("learn_build", {"level": 1, "topics": ["cash"], "hours": 1})
    new = json.dumps({"title": "نقدی کا انتظام", "objectives": ["سیکھیں"], "blocks": [{"type": "para", "text": "اردو متن"}]})
    monkeypatch.setattr(api.agent, "complete", lambda *a, **k: (new, "gemini"))
    c = api.dispatch("learn_rewrite", {"course_id": c["id"], "module": 0, "lesson": 0, "how": "urdu"})
    assert c["content"]["modules"][0]["lessons"][0]["title"] == "نقدی کا انتظام"
    content = c["content"]
    content["modules"][0]["lessons"][0]["blocks"][0]["text"] = "Edited by hand"
    c = api.dispatch("learn_edit", {"course_id": c["id"], "content": content})
    assert c["content"]["modules"][0]["lessons"][0]["blocks"][0]["text"] == "Edited by hand"
    dup = api.dispatch("learn_duplicate", {"course_id": c["id"]})
    assert dup["id"] != c["id"] and dup["content"]["title"].endswith("(copy)")


def test_assistant_makes_and_lists_courses(api):
    r = api.dispatch("agent_ask", {"question": "Make a 3 hour course on shrink for team leaders"})
    assert "Course ready" in r["answer"] and "Shrinkage" in r["answer"]
    r = api.dispatch("agent_ask", {"question": "What training courses do we have?"})
    assert "courses" in r["answer"] and "Shrinkage" in r["answer"]
