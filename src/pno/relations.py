"""2.8 Employee relations: cases with stages and deadlines, HR letters in one click, and exit interviews.

Cases follow the steps HR actually takes (show cause → reply → hearing → decision for discipline), each with a due date,
so overdue steps show on the Today list. A person's earlier warnings in the last 12 months are shown with the next step
on the ladder (counselling → verbal → written → final written → termination) so outcomes stay consistent and fair.
"""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from .db import now

KINDS = {
    "Disciplinary": [("Opened", 0), ("Show cause issued", 1), ("Reply received", 3), ("Hearing held", 7), ("Decision", 3), ("Closed", 0)],
    "Grievance": [("Opened", 0), ("Acknowledged", 2), ("Investigation", 7), ("Resolution shared", 7), ("Closed", 0)],
    "Investigation": [("Opened", 0), ("Evidence gathered", 5), ("Interviews done", 5), ("Findings", 5), ("Closed", 0)],
    "Performance improvement": [("Opened", 0), ("Plan agreed", 7), ("Mid-point review", 30), ("Final review", 30), ("Closed", 0)],
}
CATEGORIES = ["Lateness", "Absence without leave", "Misconduct", "Cash shortage", "Stock loss / theft", "Policy violation", "Harassment",
              "Health & safety", "Customer complaint", "Performance", "Pay or benefits", "Working hours / roster", "Manager behaviour", "Other"]
OUTCOMES = ["No action", "Counselling", "Verbal warning", "Written warning", "Final written warning", "Suspension", "Termination",
            "Grievance upheld", "Grievance not upheld", "PIP passed", "PIP extended", "PIP failed"]
LADDER = ["Counselling", "Verbal warning", "Written warning", "Final written warning", "Termination"]


def _j(v, d):
    try:
        return json.loads(v) if v else d
    except ValueError:
        return d


def _due(kind: str, stage: str, start: date | None = None) -> str | None:
    steps = KINDS.get(kind, KINDS["Disciplinary"])
    names = [s for s, _ in steps]
    if stage not in names or stage == "Closed":
        return None
    i = names.index(stage)
    if i + 1 >= len(steps):
        return None
    return ((start or date.today()) + timedelta(days=max(1, steps[i + 1][1]))).isoformat()


def open_case(db, kind: str, emp: str, name: str, store: str = "", section: str = "", designation: str = "", category: str = "Other",
              description: str = "", opened_on: str | None = None, confidential: bool = False) -> int:
    if kind not in KINDS:
        raise ValueError("Choose the type of case.")
    if not (emp or name):
        raise ValueError("Who is the case about?")
    od = opened_on or date.today().isoformat()
    hist = [{"at": now(), "stage": "Opened", "note": description[:500]}]
    return db.x("""INSERT INTO er_cases(kind, emp, name, store, section, designation, category, description, opened_on, stage, stage_due,
                   outcome, history, confidential, created_at, updated_at) VALUES(?,?,?,?,?,?,?,?,?,'Opened',?,NULL,?,?,?,?)""",
                (kind, str(emp or ""), name, store, section, designation, category if category in CATEGORIES else "Other", description[:4000], od,
                 _due(kind, "Opened", date.fromisoformat(od)), json.dumps(hist), int(bool(confidential)), now(), now())).lastrowid


def advance(db, case_id: int, stage: str | None = None, note: str = "", outcome: str | None = None, due: str | None = None) -> dict:
    c = db.q1("SELECT * FROM er_cases WHERE id=?", (case_id,))
    if not c:
        raise ValueError("Case not found.")
    names = [s for s, _ in KINDS[c["kind"]]]
    nxt = stage or (names[names.index(c["stage"]) + 1] if c["stage"] in names and names.index(c["stage"]) + 1 < len(names) else "Closed")
    if nxt not in names:
        raise ValueError("Unknown step.")
    if nxt == "Closed" and not (outcome or c["outcome"]) and c["kind"] != "Investigation":
        raise ValueError("Record the outcome before closing the case.")
    hist = _j(c["history"], [])
    hist.append({"at": now(), "stage": nxt, "note": note[:1000], **({"outcome": outcome} if outcome else {})})
    db.x("UPDATE er_cases SET stage=?, stage_due=?, outcome=COALESCE(?, outcome), history=?, updated_at=?, closed_on=? WHERE id=?",
         (nxt, due or _due(c["kind"], nxt), outcome if outcome in OUTCOMES else None, json.dumps(hist), now(),
          date.today().isoformat() if nxt == "Closed" else None, case_id))
    db.log("er case", f"#{case_id} → {nxt}{(' · ' + outcome) if outcome else ''}")
    return case(db, case_id)


def case(db, case_id: int) -> dict:
    c = db.q1("SELECT * FROM er_cases WHERE id=?", (case_id,))
    if not c:
        raise ValueError("Case not found.")
    d = dict(c)
    d["history"] = _j(d["history"], [])
    d["steps"] = [s for s, _ in KINDS[d["kind"]]]
    d["prior"] = prior(db, d["emp"], exclude=case_id)
    d["suggest"] = suggest_outcome(d["prior"]) if d["kind"] == "Disciplinary" else None
    d["letters"] = [dict(r) for r in db.q("SELECT id, kind, ref, created_at FROM letters WHERE case_id=? ORDER BY id", (case_id,))]
    d["overdue"] = bool(d["stage_due"] and d["stage"] != "Closed" and d["stage_due"] < date.today().isoformat())
    return d


def prior(db, emp: str, exclude: int | None = None) -> list[dict]:
    if not emp:
        return []
    since = (date.today() - timedelta(days=365)).isoformat()
    return [dict(r) for r in db.q("""SELECT id, kind, category, outcome, opened_on, stage FROM er_cases WHERE emp=? AND id!=? AND opened_on>=?
                                     ORDER BY opened_on DESC""", (emp, exclude or 0, since))]


def suggest_outcome(prev: list[dict]) -> dict:
    """The next step on the ladder given warnings in the last 12 months (guidance only; the panel decides)."""
    done = [p["outcome"] for p in prev if p.get("outcome") in LADDER]
    top = max((LADDER.index(o) for o in done), default=-1)
    nxt = LADDER[min(top + 1, len(LADDER) - 1)]
    why = (f"{len(done)} earlier warning{'s' if len(done) != 1 else ''} in 12 months, the highest being {LADDER[top].lower()}"
           if top >= 0 else "No warnings in the last 12 months")
    return {"outcome": nxt, "why": why + ". Applies only if the allegation is proven; the panel decides."}


def cases(db, status: str = "open") -> dict:
    rows = []
    t = date.today().isoformat()
    for r in db.q("SELECT * FROM er_cases ORDER BY stage='Closed', COALESCE(stage_due,'9999'), id DESC"):
        d = dict(r)
        d.pop("history", None)
        d["overdue"] = bool(d["stage_due"] and d["stage"] != "Closed" and d["stage_due"] < t)
        d["days_open"] = (date.today() - date.fromisoformat(d["opened_on"])).days if d["opened_on"] else None
        rows.append(d)
    open_ = [r for r in rows if r["stage"] != "Closed"]
    closed = [r for r in rows if r["stage"] == "Closed" and r["closed_on"]]
    dtc = [(date.fromisoformat(r["closed_on"]) - date.fromisoformat(r["opened_on"])).days for r in closed if r["opened_on"]]
    shown = rows if status == "all" else open_ if status == "open" else [r for r in rows if r["overdue"]] if status == "overdue" else closed
    return {"rows": shown, "open": len(open_), "overdue": sum(1 for r in open_ if r["overdue"]), "closed": len(closed),
            "avg_days_to_close": round(sum(dtc) / len(dtc), 1) if dtc else None,
            "by_category": Counter(r["category"] for r in rows).most_common(), "by_store": Counter(r["store"] or "—" for r in open_).most_common(10),
            "by_outcome": Counter(r["outcome"] for r in closed if r["outcome"]).most_common(),
            "kinds": list(KINDS), "categories": CATEGORIES, "outcomes": OUTCOMES}


# ---------------------------------------------------------------------------------------------- letters
LETTERS: dict[str, dict] = {
    "offer": {"title": "Offer letter", "fields": ["position", "store", "grade", "salary", "start_date", "probation_months", "reporting_to"],
              "body": ["We are pleased to offer you the position of {position} at {store}, Majid Al Futtaim Hypermarkets Pakistan (Carrefour).",
                       "Your employment will start on {start_date}, at grade {grade}, with a monthly gross salary of PKR {salary}. You will report to {reporting_to}.",
                       "The first {probation_months} months will be a probation period. Further terms are set out in your appointment letter and the company's policies.",
                       "Please confirm your acceptance by signing and returning a copy of this letter. We look forward to welcoming you to the team."]},
    "appointment": {"title": "Appointment letter", "fields": ["position", "store", "grade", "start_date", "probation_months", "reporting_to", "working_hours"],
                    "body": ["Further to your acceptance of our offer, we confirm your appointment as {position} at {store} with effect from {start_date}.",
                             "Your grade is {grade} and you will report to {reporting_to}. Normal working hours are {working_hours}, on a roster set by your manager.",
                             "You will be on probation for {probation_months} months, after which your performance will be reviewed for confirmation.",
                             "You are required to follow the company's code of conduct and all HR policies, as amended from time to time."]},
    "confirmation": {"title": "Confirmation of employment", "fields": ["position", "store", "effective_date"],
                     "body": ["We are pleased to confirm that you have successfully completed your probation period.",
                              "Your employment as {position} at {store} is confirmed with effect from {effective_date}.",
                              "Thank you for your contribution so far. We wish you continued success with Carrefour."]},
    "probation_extension": {"title": "Extension of probation", "fields": ["position", "extended_to", "areas"],
                            "body": ["Following your probation review, we have decided to extend your probation period until {extended_to}.",
                                     "During this period, we expect to see clear improvement in: {areas}. Your manager will review progress with you regularly.",
                                     "Your confirmation will be considered at the end of the extended period."]},
    "transfer": {"title": "Transfer letter", "fields": ["from_store", "to_store", "position", "effective_date", "reporting_to"],
                 "body": ["This is to inform you that you are transferred from {from_store} to {to_store} as {position}, with effect from {effective_date}.",
                          "You will report to {reporting_to}. All other terms and conditions of your employment remain unchanged.",
                          "Please complete your handover before the effective date."]},
    "promotion": {"title": "Promotion letter", "fields": ["new_position", "store", "new_grade", "effective_date", "new_salary"],
                  "body": ["Congratulations! In recognition of your performance and contribution, you are promoted to {new_position} at {store} with effect from {effective_date}.",
                           "Your new grade is {new_grade} and your revised monthly gross salary is PKR {new_salary}.",
                           "We are confident you will continue to deliver excellent results in your new role."]},
    "show_cause": {"title": "Show cause notice", "fields": ["incident_date", "allegation", "reply_by"],
                   "body": ["It has been reported that on {incident_date}: {allegation}.",
                            "This, if proven, is a breach of the company's code of conduct and policies.",
                            "You are required to explain in writing, by {reply_by}, why disciplinary action should not be taken against you. "
                            "If no reply is received by then, the matter will be decided on the information available.",
                            "This notice is issued without prejudice to any further action the company may take."]},
    "warning": {"title": "Warning letter", "fields": ["level", "incident_date", "allegation", "expected"],
                "body": ["Following the review of the incident of {incident_date} ({allegation}) and your reply, the company has decided to issue you a {level}.",
                         "You are expected to {expected}. Any repetition may lead to further disciplinary action, up to and including termination of employment.",
                         "This letter will remain on your personal file for twelve months."]},
    "experience": {"title": "Experience certificate", "fields": ["position", "store", "from_date", "to_date"],
                   "body": ["This is to certify that {name} (employee number {emp}) worked with Majid Al Futtaim Hypermarkets Pakistan (Carrefour) "
                            "from {from_date} to {to_date}, last as {position} at {store}.",
                            "During this period we found them hardworking and responsible. We wish them every success in the future.",
                            "This certificate is issued on request, without any liability on the part of the company."]},
    "relieving": {"title": "Relieving letter", "fields": ["position", "last_day", "resignation_date"],
                  "body": ["With reference to your resignation dated {resignation_date}, you are relieved from your duties as {position} at the close of business on {last_day}.",
                           "Your full and final settlement will be processed after clearance from all departments, in line with company policy.",
                           "We thank you for your services and wish you the best in your future endeavours."]},
}


def next_ref(db, kind: str) -> str:
    n = db.q1("SELECT COUNT(*) FROM letters WHERE kind=? AND substr(created_at,1,4)=?", (kind, str(date.today().year)))[0] + 1
    return f"MAF-PK/P&O/{kind.upper().replace('_', '-')}/{date.today().year}/{n:03d}"


def letter_defaults(api, kind: str, emp: str, case_id: int | None = None) -> dict:
    """The letter's fields filled from what P&O knows about the person (and the case)."""
    b = api._brain()
    p = b.people.get(str(emp))
    f = (p or {}).get("facts", {})
    g = lambda k: (f.get(k) or {}).get("value", "")  # noqa: E731
    t = date.today()
    d = {"name": (p or {}).get("name", ""), "emp": str(emp), "position": g("designation"), "store": g("store"), "grade": g("grade"),
         "reporting_to": g("manager_name"), "start_date": (t + timedelta(days=14)).strftime("%d %B %Y"), "probation_months": "3",
         "working_hours": "48 hours a week (9 hours a day with breaks)", "effective_date": t.strftime("%d %B %Y"),
         "extended_to": (t + timedelta(days=90)).strftime("%d %B %Y"), "areas": "", "from_store": g("store"), "to_store": "",
         "new_position": "", "new_grade": "", "new_salary": "", "salary": "", "incident_date": t.strftime("%d %B %Y"), "allegation": "",
         "reply_by": (t + timedelta(days=3)).strftime("%d %B %Y"), "level": "written warning", "expected": "follow the company's policies at all times",
         "from_date": "", "to_date": t.strftime("%d %B %Y"), "last_day": t.strftime("%d %B %Y"), "resignation_date": ""}
    if g("joining_date"):
        try:
            d["from_date"] = date.fromisoformat(str(g("joining_date"))[:10]).strftime("%d %B %Y")
        except ValueError:
            pass
    if p and p.get("left") and p["left"].get("date"):
        try:
            d["to_date"] = d["last_day"] = date.fromisoformat(str(p["left"]["date"])[:10]).strftime("%d %B %Y")
        except ValueError:
            pass
    if case_id:
        c = api.db.q1("SELECT * FROM er_cases WHERE id=?", (case_id,))
        if c:
            d["allegation"] = (c["description"] or c["category"] or "").strip().rstrip(".")
            d["incident_date"] = date.fromisoformat(c["opened_on"]).strftime("%d %B %Y") if c["opened_on"] else d["incident_date"]
            if c["outcome"] in ("Verbal warning", "Written warning", "Final written warning"):
                d["level"] = c["outcome"].lower()
            d["name"] = d["name"] or c["name"]
    return {"kind": kind, "title": LETTERS[kind]["title"], "fields": {k: d.get(k, "") for k in LETTERS[kind]["fields"]}, "name": d["name"], "emp": str(emp),
            "ref": next_ref(api.db, kind)}


def make_letter(db, kind: str, emp: str, name: str, fields: dict, path: str, fmt: str = "docx", case_id: int | None = None,
                signatory: str = "Director, People & Organisation") -> dict:
    if kind not in LETTERS:
        raise ValueError("Unknown letter.")
    spec = LETTERS[kind]
    missing = [k for k in spec["fields"] if not str(fields.get(k) or "").strip()]
    if missing:
        raise ValueError("Fill in: " + ", ".join(m.replace("_", " ") for m in missing))
    vals = defaultdict(str, {**fields, "name": name, "emp": emp})
    paras = [re.sub(r"\{(\w+)\}", lambda m: str(vals[m.group(1)]), t) for t in spec["body"]]
    ref = next_ref(db, kind)
    content = {"title": spec["title"], "ref": ref, "date": date.today().strftime("%d %B %Y"), "to": name, "emp": emp,
               "position": fields.get("position") or fields.get("new_position") or "", "store": fields.get("store") or fields.get("to_store") or "",
               "paras": paras, "signatory": signatory, "confidential": kind in ("show_cause", "warning")}
    (_pdf if fmt == "pdf" else _docx)(content, path)
    lid = db.x("INSERT INTO letters(kind, emp, name, ref, fields, case_id, created_at) VALUES(?,?,?,?,?,?,?)",
               (kind, emp, name, ref, json.dumps(fields), case_id, now())).lastrowid
    db.log("letter", f"{spec['title']} {ref} for {name}")
    return {"id": lid, "ref": ref, "path": path}


def _docx(c: dict, path: str) -> None:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor
    from .learning.export import ASSETS
    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Mm(22)
        s.top_margin, s.bottom_margin = Mm(16), Mm(18)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Segoe UI", Pt(10.5)
    logo = ASSETS / "maf-logo.png"
    if logo.exists():
        d.add_picture(str(logo), width=Mm(36))
    p = d.add_paragraph()
    r = p.add_run("Majid Al Futtaim Hypermarkets Pakistan · Carrefour · People & Organisation")
    r.font.size, r.font.color.rgb = Pt(8.5), RGBColor(0x76, 0x6C, 0x63)
    if c["confidential"]:
        p = d.add_paragraph()
        r = p.add_run("STRICTLY PRIVATE AND CONFIDENTIAL")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0xB3, 0x26, 0x1E)
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = f"Ref: {c['ref']}"
    t.cell(0, 1).text = f"Date: {c['date']}"
    t.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph()
    for line in [c["to"], f"Employee number: {c['emp']}" if c["emp"] else "", c["position"], c["store"]]:
        if line:
            d.add_paragraph(line).paragraph_format.space_after = Pt(0)
    d.add_paragraph()
    h = d.add_paragraph()
    r = h.add_run(c["title"].upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(13), RGBColor(0x94, 0x48, 0x1A)
    d.add_paragraph(f"Dear {c['to'].split(' ')[0] if c['to'] else 'Sir / Madam'},")
    for para in c["paras"]:
        q = d.add_paragraph(para)
        q.paragraph_format.space_after = Pt(8)
        q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_paragraph()
    d.add_paragraph("Yours sincerely,")
    d.add_paragraph("\n\n______________________________")
    d.add_paragraph(c["signatory"])
    d.add_paragraph("Majid Al Futtaim Hypermarkets Pakistan")
    if c["title"] in ("Offer letter", "Appointment letter", "Show cause notice", "Warning letter", "Transfer letter"):
        d.add_paragraph()
        d.add_paragraph("Received and acknowledged:\n\nName: ____________________   Signature: ____________________   Date: ____________")
    d.save(path)


def _pdf(c: dict, path: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    from .learning.export import ASSETS
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=16 * mm, bottomMargin=18 * mm,
                            title=c["title"], author="People & Organisation")
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=10.5, leading=15)
    small = ParagraphStyle("s", parent=base, fontSize=8.5, textColor=colors.HexColor("#766C63"))
    head = ParagraphStyle("h", parent=base, fontName="Helvetica-Bold", fontSize=13, textColor=colors.HexColor("#94481A"), spaceBefore=6, spaceAfter=8)
    body = ParagraphStyle("p", parent=base, alignment=TA_JUSTIFY, spaceAfter=8)
    story = []
    logo = ASSETS / "maf-logo.png"
    if logo.exists():
        story.append(Image(str(logo), width=36 * mm, height=18 * mm, kind="proportional", hAlign="LEFT"))
    story.append(Paragraph("Majid Al Futtaim Hypermarkets Pakistan · Carrefour · People &amp; Organisation", small))
    if c["confidential"]:
        story.append(Paragraph("<b><font color='#B3261E'>STRICTLY PRIVATE AND CONFIDENTIAL</font></b>", small))
    story.append(Spacer(1, 6 * mm))
    t = Table([[Paragraph(f"Ref: {escape(c['ref'])}", base), Paragraph(f"Date: {escape(c['date'])}", ParagraphStyle("r", parent=base, alignment=2))]],
              colWidths=[90 * mm, 76 * mm])
    t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story += [t, Spacer(1, 5 * mm)]
    for line in [c["to"], f"Employee number: {c['emp']}" if c["emp"] else "", c["position"], c["store"]]:
        if line:
            story.append(Paragraph(escape(line), base))
    story += [Spacer(1, 5 * mm), Paragraph(escape(c["title"].upper()), head),
              Paragraph(f"Dear {escape(c['to'].split(' ')[0]) if c['to'] else 'Sir / Madam'},", body)]
    story += [Paragraph(escape(x), body) for x in c["paras"]]
    story += [Spacer(1, 6 * mm), Paragraph("Yours sincerely,", base), Spacer(1, 16 * mm), Paragraph("______________________________", base),
              Paragraph(escape(c["signatory"]), base), Paragraph("Majid Al Futtaim Hypermarkets Pakistan", base)]
    if c["title"] in ("Offer letter", "Appointment letter", "Show cause notice", "Warning letter", "Transfer letter"):
        story += [Spacer(1, 10 * mm), Paragraph("Received and acknowledged: Name ____________________  Signature ____________________  Date ____________", small)]
    doc.build(story)


def letters(db) -> list[dict]:
    return [dict(r) for r in db.q("SELECT id, kind, emp, name, ref, case_id, created_at FROM letters ORDER BY id DESC LIMIT 300")]


# ---------------------------------------------------------------------------------------------- exit interviews
EXIT_REASONS = ["Better pay elsewhere", "Career growth", "Manager / supervisor", "Working hours / roster", "Workload", "Distance / transport",
                "Studies", "Family / personal", "Health", "Relocation", "Work environment", "Recognition", "Contract ended", "Other"]


def add_exit(db, emp: str, name: str, store: str = "", section: str = "", designation: str = "", leave_date: str | None = None,
             main_reason: str = "Other", reasons: list | None = None, manager_score: int | None = None, culture_score: int | None = None,
             pay_score: int | None = None, growth_score: int | None = None, would_return: bool | None = None, recommend: int | None = None,
             comments: str = "") -> int:
    if not (emp or name):
        raise ValueError("Who is leaving?")
    sc = lambda v: int(v) if v not in (None, "") and 1 <= int(v) <= 5 else None  # noqa: E731
    rec = int(recommend) if recommend not in (None, "") and 0 <= int(recommend) <= 10 else None
    return db.x("""INSERT INTO exit_interviews(emp, name, store, section, designation, leave_date, main_reason, reasons, manager_score, culture_score,
                   pay_score, growth_score, would_return, recommend, comments, created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (str(emp or ""), name, store, section, designation, leave_date, main_reason if main_reason in EXIT_REASONS else "Other",
                 json.dumps([r for r in (reasons or []) if r in EXIT_REASONS]), sc(manager_score), sc(culture_score), sc(pay_score),
                 sc(growth_score), None if would_return is None else int(bool(would_return)), rec, comments[:3000], now())).lastrowid


def exits(db, api=None) -> dict:
    rows = [dict(r) for r in db.q("SELECT * FROM exit_interviews ORDER BY COALESCE(leave_date, created_at) DESC")]
    for r in rows:
        r["reasons"] = _j(r["reasons"], [])
    n = len(rows)
    reasons = Counter(r["main_reason"] for r in rows) + Counter(x for r in rows for x in r["reasons"] if x != r["main_reason"])
    avg = lambda k: round(sum(r[k] for r in rows if r[k]) / max(1, sum(1 for r in rows if r[k])), 2) if any(r[k] for r in rows) else None  # noqa: E731
    rec = [r["recommend"] for r in rows if r["recommend"] is not None]
    enps = round(100 * (sum(1 for x in rec if x >= 9) - sum(1 for x in rec if x <= 6)) / len(rec)) if rec else None
    regrettable = []
    if api is not None:
        try:
            b = api._brain()
            for r in rows:
                p = b.people.get(str(r["emp"]))
                if not p:
                    continue
                rated = [x for x in p["ratings"] if x.get("rating") is not None]
                idx = p["leader"].score if p.get("leader") else None
                why = [f"rated {rated[-1]['rating']:g}"] if rated and rated[-1]["rating"] >= 4 else []
                if idx is not None and idx >= 70:
                    why.append(f"leader index {idx / 10:.1f}")
                if why:
                    regrettable.append({"name": r["name"], "store": r["store"], "reason": r["main_reason"], "why": " and ".join(why)})
        except Exception:  # noqa: BLE001
            pass
    return {"rows": rows, "n": n, "reasons": reasons.most_common(), "main": Counter(r["main_reason"] for r in rows).most_common(),
            "scores": {"Manager": avg("manager_score"), "Culture": avg("culture_score"), "Pay": avg("pay_score"), "Growth": avg("growth_score")},
            "would_return": round(100 * sum(1 for r in rows if r["would_return"] == 1) / max(1, sum(1 for r in rows if r["would_return"] is not None)))
            if any(r["would_return"] is not None for r in rows) else None,
            "enps": enps, "regrettable": regrettable, "by_store": Counter(r["store"] or "—" for r in rows).most_common(12), "reasons_all": EXIT_REASONS}
