"""3.0: the director's office: automatic backups, the HR calendar, the weekly brief, the store-visit checklist, the
glossary behind 'explain this number', and column names P&O teaches Smart Import.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from .paths import data_dir

KEEP_BACKUPS = 14


# ------------------------------------------------------------------------------------------- backups
def backup_dir() -> Path:
    d = data_dir() / "backups"
    d.mkdir(parents=True, exist_ok=True)
    return d


def auto_backup(db, force: bool = False) -> dict | None:
    """One copy a day, kept for two weeks. Runs at start-up; never blocks the app."""
    if getattr(db, "path", None) is None or db.get_setting("auto_backup", True) is False:
        return None
    today = date.today().isoformat()
    if not force and db.get_setting("auto_backup_last", "") == today:
        return None
    dest = backup_dir() / f"P&O auto {datetime.now():%Y-%m-%d %H%M}.sqlite3"
    db.backup(dest)
    db.set_setting("auto_backup_last", today)
    files = sorted(backup_dir().glob("P&O auto *.sqlite3"))
    for old in files[:-KEEP_BACKUPS]:
        try:
            old.unlink()
        except OSError:
            pass
    return {"path": str(dest)}


def backups() -> list[dict]:
    out = []
    for f in sorted(backup_dir().glob("*.sqlite3"), reverse=True):
        st = f.stat()
        out.append({"name": f.name, "path": str(f), "size_kb": round(st.st_size / 1024), "at": datetime.fromtimestamp(st.st_mtime).isoformat(timespec="minutes")})
    return out


# ------------------------------------------------------------------------------------------- synonyms
def synonyms(db) -> dict:
    from .smart import schema
    extra = db.get_setting("synonyms", {}) or {}
    return {"fields": [{"field": k, "label": v["label"], "builtin": schema._BASE_SYN[k][:12], "extra": extra.get(k, [])} for k, v in schema.FIELDS.items()]}


def save_synonyms(db, field: str, words: list[str]) -> dict:
    from .smart import schema
    if field not in schema.FIELDS:
        raise ValueError("Unknown field.")
    extra = db.get_setting("synonyms", {}) or {}
    clean = sorted({w.strip().casefold() for w in words if w and w.strip()})[:40]
    if clean:
        extra[field] = clean
    else:
        extra.pop(field, None)
    db.set_setting("synonyms", extra)
    schema.set_user_synonyms(extra)
    return synonyms(db)


def load_synonyms(db) -> None:
    from .smart import schema
    schema.set_user_synonyms(db.get_setting("synonyms", {}) or {})


# ------------------------------------------------------------------------------------------- HR calendar
def calendar(api, month: str | None = None) -> dict:
    """Every dated HR event in a month: anniversaries, probation ends, leave, ER steps, goals, actions, follow-ups, holidays."""
    m = month or date.today().strftime("%Y-%m")
    y, mo = int(m[:4]), int(m[5:7])
    start = date(y, mo, 1)
    end = (date(y + mo // 12, mo % 12 + 1, 1) - timedelta(days=1))
    s, e = start.isoformat(), end.isoformat()
    ev: dict[str, list] = defaultdict(list)
    add = lambda d, kind, text, go="", emp="": ev[str(d)[:10]].append({"kind": kind, "text": text, "go": go, "emp": emp})  # noqa: E731
    db = api.db
    try:
        b = api._brain()
        for p in b.people.values():
            if p["left"]:
                continue
            jd = (p["facts"].get("joining_date") or {}).get("value")
            try:
                j = date.fromisoformat(str(jd)[:10])
            except (TypeError, ValueError):
                continue
            if j.month == mo and y > j.year:
                try:
                    d = j.replace(year=y)
                except ValueError:
                    d = date(y, 3, 1)
                yrs = y - j.year
                if yrs in (1, 3, 5, 10, 15, 20, 25, 30) or canon_rank(p) >= 3:
                    add(d, "anniversary", f"{p['name']}: {yrs} year{'s' if yrs > 1 else ''}", "brain", p["emp"])
        from . import org_health
        for r in org_health.probation(api._po(), 3)["rows"]:
            if s <= r["ends"] <= e:
                add(r["ends"], "probation", f"Probation ends: {r['name']} ({r['store'] or ''})", "organisation")
    except Exception:  # noqa: BLE001
        pass
    for r in db.q("SELECT emp_no, d_from, d_to, type FROM leaves WHERE d_from <= ? AND d_to >= ?", (e, s)):
        add(max(r["d_from"], s), "leave", f"{r['type']}: {r['emp_no']} until {r['d_to']}", "attendance", r["emp_no"])
    for r in db.q("SELECT id, kind, name, stage, stage_due FROM er_cases WHERE stage!='Closed' AND stage_due BETWEEN ? AND ?", (s, e)):
        add(r["stage_due"], "er", f"{r['kind']} #{r['id']} {r['name']}: next step after '{r['stage']}'", "relations")
    for r in db.q("SELECT title, due FROM goals WHERE status='open' AND due BETWEEN ? AND ?", (s, e)):
        add(r["due"], "goal", f"Goal due: {r['title']}", "goals")
    for r in db.q("SELECT title, owner, due FROM actions WHERE status IN ('open','in progress') AND due BETWEEN ? AND ?", (s, e)):
        add(r["due"], "action", f"{r['title']}{' · ' + r['owner'] if r['owner'] else ''}", "actions")
    for r in db.q("SELECT subject, follow_up, follow_due FROM meetings WHERE follow_due BETWEEN ? AND ? AND follow_up != ''", (s, e)):
        add(r["follow_due"], "followup", f"Follow up with {r['subject']}: {r['follow_up']}", "notes")
    for r in db.q("SELECT day, name FROM holidays WHERE day BETWEEN ? AND ?", (s, e)):
        add(r["day"], "holiday", r["name"])
    try:
        for r in db.q("SELECT title, due FROM assignments a JOIN courses c ON c.id=a.course_id WHERE a.status!='completed' AND a.due BETWEEN ? AND ? "
                      "GROUP BY c.id, a.due", (s, e)):
            add(r["due"], "training", f"Training due: {r['title']}", "learning")
    except Exception:  # noqa: BLE001
        pass
    days = []
    d = start
    while d <= end:
        days.append({"day": d.isoformat(), "events": ev.get(d.isoformat(), [])})
        d += timedelta(days=1)
    counts = defaultdict(int)
    for x in ev.values():
        for i in x:
            counts[i["kind"]] += 1
    return {"month": m, "label": start.strftime("%B %Y"), "first_weekday": start.weekday(), "days": days, "counts": dict(counts)}


def canon_rank(p: dict) -> int:
    from .po import canon
    return canon.role_rank((p["facts"].get("designation") or {}).get("value", ""))


# ------------------------------------------------------------------------------------------- weekly brief
def weekly_brief(api) -> dict:
    """One page for Monday morning: the numbers, what changed, what needs the director, and what's coming."""
    from . import actions, people_insights
    sections = []
    ins = api._po()
    if not ins.empty:
        h = ins.headline()
        line = []
        if h.get("sales_vs_budget_pct") is not None:
            line.append(f"sales at {h['sales_vs_budget_pct']:.1f}% of budget")
        if h.get("productivity_vs_target_pct") is not None:
            line.append(f"productivity {h['productivity_vs_target_pct']:+.1f}% vs target")
        line.append(f"{h['stores_on_track']} of {h['stores']} stores on track, {h['stores_at_risk']} at risk")
        if h.get("leavers_12m") is not None:
            line.append(f"{h['leavers_12m']} leavers in 12 months ({h.get('voluntary_pct') or 0:.0f}% voluntary)")
        sections.append({"title": "The numbers", "items": ["Country: " + ", ".join(line) + "."]})
        alerts = [a for a in ins.alerts(30) if a["severity"] == "High"][:5]
        if alerts:
            sections.append({"title": "Stores that need attention", "items": [f"{a['store']}: {a['message']}" for a in alerts]})
        ranked = sorted([l for l in ins.leaders() if l.score is not None], key=lambda l: -l.score)
        if ranked:
            sections.append({"title": "Leaders", "items": [f"Top: {', '.join(f'{l.name} ({l.store}, {l.score / 10:.1f})' for l in ranked[:3])}.",
                                                            f"Needs support: {', '.join(f'{l.name} ({l.store}, {l.score / 10:.1f})' for l in ranked[-3:])}."]})
    today = actions.today(api)["items"]
    if today:
        sections.append({"title": "Needs you this week", "items": [f"{i['area']}: {i['text']}" for i in today[:8]]})
    try:
        pi = people_insights.report(api)
        items = []
        a = pi.get("absence") or {}
        if a.get("available"):
            items.append(f"Absence {a['average']}% of scheduled days" + (f"; {a['peak']['day']} is {a['peak']['x']}× the average" if a.get("peak") else "") +
                         (f"; {a['bradford_high']} people with a repeat-absence pattern" if a.get("bradford_high") else "") + ".")
        w = pi.get("wellbeing") or {}
        if w.get("n"):
            items.append(f"{w['n']} people on very long hours or without rest.")
        t = pi.get("attrition") or {}
        if t.get("regrettable_n"):
            items.append(f"{t['regrettable_n']} strong performers have left ({t['regrettable_voluntary']} resigned).")
        for d in (pi.get("drivers") or {}).get("drivers", [])[:2]:
            items.append(d["text"])
        wk = pi.get("week") or {}
        if wk.get("anniversaries"):
            items.append("Anniversaries: " + ", ".join(f"{x['name']} ({x['years']})" for x in wk["anniversaries"][:5]) + ".")
        if items:
            sections.append({"title": "People", "items": items})
    except Exception:  # noqa: BLE001
        pass
    nxt = date.today() + timedelta(days=7)
    cal = calendar(api)
    coming = [f"{x['day'][8:]}: {i['text']}" for x in cal["days"] if date.today().isoformat() <= x["day"] <= nxt.isoformat() for i in x["events"]
              if i["kind"] in ("er", "probation", "goal", "holiday", "training")][:8]
    if coming:
        sections.append({"title": "Coming up", "items": coming})
    text = "\n\n".join(f"{s['title'].upper()}\n" + "\n".join(f"• {i}" for i in s["items"]) for s in sections)
    return {"title": f"P&O weekly brief · week of {date.today() - timedelta(days=date.today().weekday()):%d %b %Y}", "sections": sections, "text": text}


# ------------------------------------------------------------------------------------------- store visit
CHECKLIST = {
    "People & rosters": ["Roster posted and matches the people on the floor", "Breaks staggered, no section left uncovered",
                         "New joiners have a buddy and know their trainer", "Absent people today were followed up"],
    "Leadership": ["Morning briefing held and targets shared", "Section managers know their sales, waste and productivity",
                   "Open actions from the last visit are closed"],
    "Customer & floor": ["Shelves full, price labels right, no gaps in top lines", "Fresh counters staffed at peak hours",
                         "Checkouts open match the queue"],
    "Compliance & wellbeing": ["Uniform and grooming standard", "Attendance punches match the roster (no manual fixes needed)",
                               "Notice board up to date (policies, holidays, recognition)", "Staff room clean; water and rest area OK"],
    "Talk to people": ["Spoke to at least 3 staff members", "Any concern raised is written down with an owner"],
}


def visit_template(api, store: str) -> dict:
    """The checklist with this store's own facts to check (at-risk KPIs, open cases, leaders to see)."""
    facts = []
    ins = api._po()
    if not store:
        names = ins.stores() if not ins.empty else []
        if not names:
            names = sorted({r["name"] for r in api.db.q("SELECT name FROM stores")})
        return {"stores": names}
    try:
        row = next((r for r in ins.store_scorecard() if r["store"] == store), None)
        if row:
            if row.get("sales_vs_budget_pct") is not None:
                facts.append(f"Sales at {row['sales_vs_budget_pct']:.1f}% of budget")
            if row.get("productivity_vs_target_pct") is not None:
                facts.append(f"Productivity {row['productivity_vs_target_pct']:+.1f}% vs target")
            if row.get("shrink_worst_section"):
                facts.append(f"Worst shrinkage: {row['shrink_worst_section']}")
            if row.get("leavers_12m"):
                facts.append(f"{row['leavers_12m']} leavers in 12 months")
        low = sorted([l for l in ins.leaders() if l.store == store and l.score is not None], key=lambda l: l.score)[:3]
        if low:
            facts.append("Spend time with: " + ", ".join(f"{l.name} ({l.section_text})" for l in low))
        cases = api.db.q("SELECT COUNT(*) n FROM er_cases WHERE stage!='Closed' AND store=?", (store,))[0]["n"]
        if cases:
            facts.append(f"{cases} open employee-relations case{'s' if cases > 1 else ''}")
    except Exception:  # noqa: BLE001
        pass
    return {"store": store, "facts": facts, "sections": CHECKLIST}


def save_visit(api, store: str, answers: dict, notes: str = "", actions_: list[str] | None = None) -> dict:
    from . import actions, workplace
    total = sum(len(v) for v in CHECKLIST.values())
    ok = sum(1 for v in answers.values() if v == "ok")
    issues = [k for k, v in answers.items() if v == "issue"]
    score = round(100 * ok / max(1, sum(1 for v in answers.values() if v in ("ok", "issue")))) if answers else None
    body = f"Store visit checklist: {ok}/{total} OK" + (f", score {score}%" if score is not None else "") + ".\n"
    if issues:
        body += "Issues: " + "; ".join(issues) + ".\n"
    if notes.strip():
        body += notes.strip()
    nid = workplace.add_note(api.db, store, body, kind="Store visit", store=store, subject_key=f"store:{store}")
    made = []
    for a in (actions_ or []):
        if a.strip():
            made.append(actions.add(api.db, a.strip(), store=store, area="Store visit", source="Store visit"))
    api.db.set_setting(f"visit:{store}:{date.today().isoformat()}", {"answers": answers, "score": score})
    return {"note_id": nid, "score": score, "issues": len(issues), "actions": len(made)}


# ------------------------------------------------------------------------------------------- explain this number
GLOSSARY: dict[str, dict[str, str]] = {
    "leader index": {"what": "A 0-10 score for each leader, combining sales vs budget, productivity vs target, shrinkage, stock days, "
                     "performance rating, team retention and customer NPS, each weighted (Leaders & talent → Weights).",
                     "source": "Sales, productivity, shrinkage, rating, turnover and NPS sheets", "good": "7+ strong · 5-7 solid · below 5 needs support"},
    "nps": {"what": "Net Promoter Score: % of customers who would recommend us (9-10) minus % who would not (0-6). Ranges -100 to +100.",
            "source": "Customer NPS sheet, by city and category", "good": "50+ excellent · 30-50 good · below 0 a warning"},
    "sales vs budget": {"what": "Actual net sales as a % of budget for the period selected (100% = on plan).", "source": "Store sales report (BO 200-10-05)",
                        "good": "100% or more of budget is on plan"},
    "productivity": {"what": "Sales per paid hour against the target set for the store or section.", "source": "Section productivity report",
                     "good": "at or above target"},
    "shrinkage": {"what": "Stock lost (damage, expiry, theft, errors) as % of sales, against its budget. Variance = actual − budget.",
                  "source": "Shrinkage & stock days sheet", "good": "variance 0 or positive"},
    "stock days": {"what": "How many days of sales the stock on hand would cover, against budget.", "source": "Shrinkage & stock days sheet",
                   "good": "at or below budget"},
    "turnover": {"what": "Leavers in the last 12 months ÷ current headcount × 100.", "source": "Turnover sheet and Basic Sheet",
                 "good": "lower is better; watch voluntary leavers"},
    "absence": {"what": "Absent days ÷ scheduled days (worked + absent), after leave, holidays and the weekly off are taken out.",
                "source": "Biometric attendance and leave register", "good": "below 3% good · above 5% a concern"},
    "bradford": {"what": "Bradford factor = (number of absence spells)² × total days absent. Many short absences score higher than one long one.",
                 "source": "Biometric attendance", "good": "under 50 normal · 125+ talk · 250+ formal review"},
    "presence": {"what": "Days worked ÷ days due (worked + unexplained absence).", "source": "Biometric attendance", "good": "95%+"},
    "rating": {"what": "The latest performance rating, 1-5, from the performance cycle.", "source": "Performance rating sheet",
               "good": "4-5 strong · 3 solid · 1-2 underperforming"},
    "enps": {"what": "Employee Net Promoter Score: % who would recommend us as an employer (9-10) minus % who would not (0-6).",
             "source": "Engagement surveys and exit interviews", "good": "above 0 positive · 20+ strong"},
    "data health": {"what": "How complete and fresh the imported data is for the month: each expected report present, dated and matched.",
                    "source": "Import history", "good": "85+ reliable"},
}


def explain(term: str) -> dict | None:
    t = (term or "").casefold()
    if "pattern" in t or "bradford" in t:
        return {"term": "bradford", **GLOSSARY["bradford"]}
    for k, v in sorted(GLOSSARY.items(), key=lambda kv: -len(kv[0])):
        if k in t or t in k:
            return {"term": k, **v}
    aliases = {"index": "leader index", "score": "leader index", "leavers": "turnover", "retention": "turnover", "attendance": "presence",
               "sales": "sales vs budget", "shrink": "shrinkage", "pattern absence": "bradford", "women": None}
    for a, k in aliases.items():
        if a in t and k:
            return {"term": k, **GLOSSARY[k]}
    return None


def brief_file(api, path: str, fmt: str = "pdf") -> str:
    """The weekly brief as a branded one-pager (PDF or Word)."""
    b = weekly_brief(api)
    from .learning.export import ASSETS
    logo = ASSETS / "maf-logo.png"
    if fmt == "docx":
        import docx
        from docx.shared import Mm, Pt, RGBColor
        d = docx.Document()
        for s in d.sections:
            s.left_margin = s.right_margin = Mm(20)
        st = d.styles["Normal"]
        st.font.name, st.font.size = "Calibri", Pt(10.5)
        if logo.exists():
            d.add_picture(str(logo), width=Mm(34))
        h = d.add_paragraph()
        r = h.add_run(b["title"])
        r.bold, r.font.size, r.font.color.rgb = True, Pt(16), RGBColor(0x94, 0x48, 0x1A)
        for sec in b["sections"]:
            p = d.add_paragraph()
            rr = p.add_run(sec["title"].upper())
            rr.bold, rr.font.size, rr.font.color.rgb = True, Pt(11), RGBColor(0xB8, 0x86, 0x0B)
            for i in sec["items"]:
                d.add_paragraph(i, style="List Bullet")
        d.add_paragraph().add_run("Confidential · People & Organisation · Majid Al Futtaim Hypermarkets Pakistan").font.size = Pt(8)
        d.save(path)
        return path
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=14 * mm, bottomMargin=16 * mm, title=b["title"])
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=10, leading=14.5, textColor=colors.HexColor("#2A1A11"))
    title = ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=colors.white)
    head = ParagraphStyle("h", parent=base, fontName="Helvetica-Bold", fontSize=11, textColor=colors.HexColor("#94481A"), spaceBefore=10, spaceAfter=4)
    small = ParagraphStyle("s", parent=base, fontSize=8, textColor=colors.HexColor("#766C63"))
    band = Table([[Paragraph(escape(b["title"]), title)]], colWidths=[174 * mm])
    band.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#94481A")), ("LEFTPADDING", (0, 0), (-1, -1), 10),
                              ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 11)]))
    story = []
    if logo.exists():
        story += [Image(str(logo), width=30 * mm, height=15 * mm, kind="proportional", hAlign="LEFT"), Spacer(1, 3 * mm)]
    story += [band, Spacer(1, 4 * mm)]
    accents = ["#B8860B", "#0EA5E9", "#DC2626", "#16A34A", "#7A5C9E", "#C2410C"]
    for k, sec in enumerate(b["sections"]):
        head.textColor = colors.HexColor(accents[k % len(accents)])
        story.append(Paragraph(escape(sec["title"].upper()), ParagraphStyle(f"h{k}", parent=head)))
        story.append(ListFlowable([ListItem(Paragraph(escape(i), base), leftIndent=10) for i in sec["items"]], bulletType="bullet", start="•", leftIndent=10))
    story += [Spacer(1, 6 * mm), Paragraph("Confidential · People &amp; Organisation · Majid Al Futtaim Hypermarkets Pakistan (Carrefour)", small)]
    doc.build(story)
    return path
