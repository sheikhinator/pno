"""One-click import of the MAF P&O review workbook (all tabs at once).

The workbook mixes ten very different layouts: block-per-store tables, a
side-by-side store matrix, a pivot table, a 31-day roster grid and the
periodic store performance export. Each tab is recognised by its header
signature (not its sheet name, which changes month to month) and parsed by
a dedicated reader into canonical records.

Personal contact fields (birth date, phone, e-mail, marital status) are
never read into the database.
"""

from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

from . import canon

PACK_KINDS: dict[str, str] = {
    "po_master": "Employee master (Basic Sheet)",
    "po_role": "Time in current role",
    "po_rating": "Performance ratings",
    "po_shift": "Shift roster",
    "po_sales": "Store performance (sales)",
    "po_productivity": "Section productivity",
    "po_shrink": "Shrinkage & stock days",
    "po_nps": "Customer NPS",
    "po_turnover": "Turnover (leavers)",
    "po_da": "Disciplinary actions",
}

PII_SKIPPED = ("Birth Date", "Age", "Local Email", "Local Mobile Number", "Marital Status", "HR Business Email")


class Pct(float):
    """A spreadsheet number that was formatted as a percentage (already ×100)."""


@dataclass
class PackSheet:
    kind: str
    sheet: str
    records: list[dict[str, Any]]
    warnings: list[str] = field(default_factory=list)
    summary: str = ""

    @property
    def label(self) -> str:
        return PACK_KINDS.get(self.kind, self.kind)


@dataclass
class PackPlan:
    path: str
    checksum: str
    sheets: list[PackSheet]
    skipped: list[str]
    period_hint: str | None = None


# ------------------------------------------------------------- reading --

def _read_xlsx(path: Path) -> list[tuple[str, list[list[Any]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[list[Any]] = []
            for row in worksheet.iter_rows():
                values = []
                for cell in row:
                    value = cell.value
                    fmt = getattr(cell, "number_format", "") or ""
                    if isinstance(value, (int, float)) and not isinstance(value, bool) and "%" in fmt:
                        value = Pct(float(value) * 100)
                    values.append(value)
                while values and (values[-1] is None or str(values[-1]).strip() == ""):
                    values.pop()
                rows.append(values)
            sheets.append((worksheet.title, rows))
    finally:
        workbook.close()
    return sheets


def _read_xls(path: Path) -> list[tuple[str, list[list[Any]]]]:
    import xlrd

    book = xlrd.open_workbook(str(path), formatting_info=False)
    sheets = []
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            values = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    value = None
                values.append(value)
            rows.append(values)
        sheets.append((sheet.name, rows))
    return sheets


def _read_text(path: Path) -> list[tuple[str, list[list[Any]]]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            content = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    delimiter = "\t" if content.count("\t") >= content.count(",") else ","
    rows = [list(row) for row in csv.reader(content.splitlines(), delimiter=delimiter)]
    return [(path.stem, rows)]


def read_sheets(path: str | Path) -> list[tuple[str, list[list[Any]]]]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(source)
    if suffix == ".xls":
        return _read_xls(source)
    return _read_text(source)


# ------------------------------------------------------------- values --

def _t(value: Any) -> str:
    text = canon.text(value)
    return "" if text.startswith("#N/A") or text in {"#REF!", "#VALUE!"} else text


def num(value: Any) -> float | None:
    """Parse 1,234 / (4.1%) / ' -   ' / '▲ 3%' style values; percent text → points."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip().replace(" ", "")
    raw = re.sub(r"^[▲▼▬]\s*", "", raw)
    if raw in {"", "-", "—", "#N/A", "#DIV/0!", "N/A", "Unrated"}:
        return None
    negative = raw.startswith("(") and raw.endswith(")")
    raw = raw.strip("()").replace(",", "").replace("%", "").strip()
    try:
        number = float(raw)
    except ValueError:
        return None
    return -number if negative else number


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _t(value)
    if not raw:
        return None
    for fmt in ("%d/%m/%Y", "%d-%b-%y", "%d-%b-%Y", "%d-%B-%Y", "%d %b %Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def _norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9%&]+", " ", _t(value).casefold()).strip()


def _find_header(rows: list[list[Any]], *needles: str, limit: int = 40) -> int | None:
    wanted = [_norm(n) for n in needles]
    for index, row in enumerate(rows[:limit]):
        cells = [_norm(c) for c in row]
        if all(any(w == c or (len(w) > 4 and w in c) for c in cells) for w in wanted):
            return index
    return None


def _column(header: list[Any], *names: str) -> int | None:
    normalised = [_norm(h) for h in header]
    for name in names:
        target = _norm(name)
        for index, value in enumerate(normalised):
            if value == target:
                return index
    for name in names:
        target = _norm(name)
        for index, value in enumerate(normalised):
            if target and target in value:
                return index
    return None


def _cell(row: list[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


# ------------------------------------------------------------ parsers --

def parse_productivity(rows: list[list[Any]], sheet: str) -> PackSheet:
    records, warnings = [], []
    header_index = _find_header(rows, "Store Name", "Store Productivity")
    header = rows[header_index]
    cols = {
        "store": _column(header, "Store Name"),
        "section": _column(header, "Dept/Section"),
        "actual": _column(header, "Store Productivity"),
        "ly": _column(header, "Store Productivity LY"),
        "target": _column(header, "Target Productivity"),
        "forecast": _column(header, "Forecasted Productivity"),
    }
    for row in rows[header_index + 1:]:
        store_raw = _t(_cell(row, cols["store"]))
        if not store_raw or _norm(store_raw) == "store name":
            continue
        section_raw = _t(_cell(row, cols["section"]))
        actual = num(_cell(row, cols["actual"]))
        if actual is None:
            continue
        records.append({
            "store": canon.store(store_raw),
            "store_raw": store_raw,
            "section": canon.section(section_raw),
            "section_raw": section_raw,
            "productivity": actual,
            "productivity_ly": num(_cell(row, cols["ly"])),
            "target": num(_cell(row, cols["target"])),
            "forecast": num(_cell(row, cols["forecast"])),
            "scope": sheet,
        })
    missing_forecast = sum(1 for r in records if r["forecast"] is None)
    if missing_forecast:
        warnings.append(f"{missing_forecast} section(s) have no forecast; forecast-vs-target is left blank, not −100%.")
    stores = len({r["store"] for r in records})
    return PackSheet("po_productivity", sheet, records, warnings, f"{len(records)} section rows · {stores} stores")


def parse_roster(rows: list[list[Any]], sheet: str) -> PackSheet:
    header_index = _find_header(rows, "Employee Number", "Employee Name")
    header = rows[header_index]
    cols = {
        "store": _column(header, "Business Unit"),
        "department": _column(header, "Department"),
        "section": _column(header, "Section"),
        "id": _column(header, "Employee Number"),
        "name": _column(header, "Employee Name"),
        "designation": _column(header, "Designation"),
    }
    day_cols = [(i, parse_date(h)) for i, h in enumerate(header)]
    day_cols = [(i, d) for i, d in day_cols if d is not None]
    records, warnings = [], []
    people = set()
    repaired = 0
    unknown_codes: set[str] = set()
    for row in rows[header_index + 1:]:
        emp = canon.employee_id(_cell(row, cols["id"]))
        if not emp:
            continue
        people.add(emp)
        raw_name = _t(_cell(row, cols["name"]))
        name = canon.repair_name(raw_name)
        repaired += name != raw_name
        base = {
            "store": canon.store(_cell(row, cols["store"])),
            "department": _t(_cell(row, cols["department"])),
            "section": canon.section(_cell(row, cols["section"])),
            "section_raw": _t(_cell(row, cols["section"])),
            "employee_id": emp,
            "employee_name": name,
            "designation": _t(_cell(row, cols["designation"])),
        }
        for index, day in day_cols:
            code = _t(_cell(row, index))
            shift = canon.shift(code)
            if shift == "Other":
                unknown_codes.add(code)
            records.append({**base, "date": day.isoformat(), "shift": shift, "shift_raw": code})
    if repaired:
        warnings.append(f"Repaired {repaired} garbled employee name(s) (e.g. 'MuhMORmad' → 'Muhammad').")
    if unknown_codes:
        warnings.append("Unrecognised shift codes kept as 'Other': " + ", ".join(sorted(unknown_codes)))
    span = f"{day_cols[0][1]:%d %b} – {day_cols[-1][1]:%d %b %Y}" if day_cols else "no dates"
    return PackSheet("po_shift", sheet, records, warnings, f"{len(people)} people · {len(day_cols)} days ({span})")


def parse_rating(rows: list[list[Any]], sheet: str) -> PackSheet:
    header_index = _find_header(rows, "Subject Employee ID")
    header = rows[header_index]
    c = lambda *n: _column(header, *n)  # noqa: E731
    cols = {
        "status": c("User Status"), "form": c("Form Template Name"), "id": c("Subject Employee ID"),
        "name": c("Subject Full Name"), "store": c("Subject Division/Entity"), "position": c("Subject Position Title"),
        "dept": c("Subject Department"), "grade": c("Subject Pay Grade"), "manager": c("Manager Full Name"),
        "rating": c("Overall Performance Rating"), "label": c("Overall Performance Rating Description"),
        "calc": c("Unadjusted Calculated Overall Performance Rating"), "section_rating": c("Section Rating"),
    }
    records = []
    for row in rows[header_index + 1:]:
        emp = canon.employee_id(_cell(row, cols["id"]))
        if not emp:
            continue
        form = _t(_cell(row, cols["form"]))
        year = re.search(r"(20\d\d)", form)
        grade_raw = _t(_cell(row, cols["grade"]))
        grade = re.match(r"^[A-Z0-9]+-(\d+)", grade_raw)
        records.append({
            "employee_id": emp,
            "employee_name": canon.repair_name(_cell(row, cols["name"])),
            "store": canon.store(re.sub(r"\(MAFR-[^)]*\)", "", _t(_cell(row, cols["store"])))),
            "designation": _t(_cell(row, cols["position"])),
            "section": canon.section(re.sub(r"\(MAFR-[^)]*\)", "", _t(_cell(row, cols["dept"])))),
            "grade": (grade.group(1).lstrip("0") or grade.group(1)) if grade else grade_raw,
            "manager_name": canon.repair_name(_cell(row, cols["manager"])),
            "cycle": int(year.group(1)) if year else None,
            "form": form,
            "role_change": "role change" in form.casefold(),
            "rating": num(_cell(row, cols["rating"])),
            "rating_label": _t(_cell(row, cols["label"])) or "Unrated",
            "calculated": num(_cell(row, cols["calc"])),
            "section_rating": num(_cell(row, cols["section_rating"])),
            "status": _t(_cell(row, cols["status"])),
        })
    unrated = sum(1 for r in records if r["rating"] is None)
    warnings = [f"{unrated} form(s) are Unrated and are excluded from rating averages."] if unrated else []
    cycles = sorted({r["cycle"] for r in records if r["cycle"]})
    people = len({r["employee_id"] for r in records})
    return PackSheet("po_rating", sheet, records, warnings, f"{len(records)} forms · {people} people · cycles {', '.join(map(str, cycles))}")


_SALES_FIELDS = [
    "actual", "forecast", "budget", "growth_pct", "var_forecast_pct", "var_budget_pct",
    "weight_store_pct", "weight_country_pct", "customers", "customers_growth_pct", "penetration_pct",
    "items", "items_growth_pct", "basket", "basket_growth_pct", "asp", "asp_growth_pct",
    "margin_pct", "waste_pct", "margin_after_waste_pct", "avg_stock", "oos_pct",
]
_SUBTOTALS = {
    "cgd fmg": "Consumer Goods (FMG)", "cgd opss": "OPSS", "consumer goods": "Consumer Goods",
    "fresh food": "Fresh Food", "food": "Food", "light house hold": "Light Household",
    "heavy house hold": "Heavy Household", "non food": "Non Food", "textile": "Textile",
    "consignment": "Consignment", "total hypermarket": "Total", "total store": "Total Store",
}


def parse_sales(rows: list[list[Any]], sheet: str) -> PackSheet:
    records, warnings = [], []
    store_name = ""
    store_raw = ""
    period_from = period_to = None
    name_col = 0
    in_table = False
    stores_seen: list[str] = []
    for row in rows:
        cells = [_t(c) for c in row]
        joined = " | ".join(c for c in cells if c)
        if not joined:
            continue
        store_match = re.search(r"Store Name:\s*([^|]+)", joined)
        if store_match:
            store_raw = store_match.group(1).strip()
            store_name = canon.store(store_raw)
            stores_seen.append(store_name)
            in_table = False
            period = re.search(r"Report Period\s*:\s*(\S+)\s*-\s*(\S+)", joined)
            if period:
                period_from, period_to = parse_date(period.group(1)), parse_date(period.group(2))
            continue
        if "Section Code Name" in cells:
            name_col = cells.index("Section Code Name")
            in_table = True
            continue
        if not in_table or not store_name:
            continue
        label = cells[name_col] if name_col < len(cells) else ""
        if not label or label.casefold() == "z" or label.startswith("Report Name"):
            continue
        values = list(row[name_col + 1:name_col + 1 + len(_SALES_FIELDS)])
        metrics = {f: num(_cell(values, i)) for i, f in enumerate(_SALES_FIELDS)}
        if metrics["actual"] is None:
            continue
        code = re.match(r"^([SX]\d{3})\s*-\s*(.+)$", label)
        if code:
            level = "section"
            line = code.group(2).strip()
            section = canon.SALES_CODE.get(code.group(1), "Consignment" if label.startswith("X") else canon.section(line))
        else:
            key = _norm(label)
            level = "total" if key in {"total hypermarket", "total store"} else "subtotal"
            line = _SUBTOTALS.get(key, label)
            section = ""
        growth = metrics["growth_pct"]
        last_year = None
        if growth is not None and growth > -100 and metrics["actual"]:
            last_year = metrics["actual"] / (1 + growth / 100)
        records.append({
            "store": store_name, "store_raw": store_raw, "line": line, "code": code.group(1) if code else "",
            "section": section, "level": level, "last_year": last_year,
            "period_from": period_from.isoformat() if period_from else None,
            "period_to": period_to.isoformat() if period_to else None,
            **metrics,
        })
    if not records:
        warnings.append("No store blocks were recognised in this report.")
    duplicates = len(stores_seen) - len(set(stores_seen))
    if duplicates:
        warnings.append(f"{duplicates} store block(s) appear more than once; later blocks are kept as separate rows.")
    span = f"{period_from:%d %b} – {period_to:%d %b %Y}" if period_from and period_to else "period not stated"
    return PackSheet("po_sales", sheet, records, warnings, f"{len(set(stores_seen))} stores · {span}")


def parse_nps(rows: list[list[Any]], sheet: str) -> PackSheet:
    header_index = _find_header(rows, "Details", "NPS")
    header = rows[header_index]
    cols = {"details": _column(header, "Details"), "section": _column(header, "Section"),
            "category": _column(header, "Product Influence"), "nps": _column(header, "NPS")}
    records = []
    for row in rows[header_index + 1:]:
        details = _t(_cell(row, cols["details"]))
        if not details:
            continue
        year = re.search(r"(20\d\d)", details)
        section_raw = _t(_cell(row, cols["section"]))
        category = _t(_cell(row, cols["category"]))
        section = canon.section(section_raw)
        by_category = canon.section(category)
        if section in {"Deli", "", section_raw} and by_category in canon.DEPARTMENT_OF:
            section = by_category
        records.append({
            "city": canon.city(details) or details.split(" ")[0],
            "year": int(year.group(1)) if year else None,
            "section": section, "section_raw": section_raw, "category": category,
            "nps": num(_cell(row, cols["nps"])),
        })
    blank = sum(1 for r in records if r["nps"] is None)
    warnings = [f"{blank} categories have no score ('-') and are excluded from averages."] if blank else []
    cities = sorted({r["city"] for r in records})
    years = sorted({r["year"] for r in records if r["year"]})
    return PackSheet("po_nps", sheet, records, warnings, f"{len(records)} scores · {', '.join(cities)} · {', '.join(map(str, years))}")


def parse_shrink(rows: list[list[Any]], sheet: str) -> PackSheet:
    store_row = _find_header(rows, "Sections")
    label_col = _column(rows[store_row], "Sections") or 0
    stores = [(i, _t(v)) for i, v in enumerate(rows[store_row]) if _t(v) and _norm(v) != "sections"]
    metric_row = store_row + 2
    records, warnings = [], []
    for row in rows[metric_row + 1:]:
        label = _t(_cell(row, label_col))
        if not label:
            continue
        for start, raw_store in stores:
            values = [_cell(row, start + k) for k in range(5)]
            budget, actual, variance, days_budget, days_actual = (num(v) for v in values)
            if all(v is None for v in (budget, actual, days_budget, days_actual)):
                continue
            records.append({
                "store": canon.store(raw_store), "store_raw": raw_store,
                "section": canon.section(label), "section_raw": label,
                "shrink_budget_pct": budget, "shrink_actual_pct": actual,
                "shrink_variance_pct": variance if variance is not None else (
                    actual - budget if actual is not None and budget is not None else None),
                "stock_days_budget": days_budget, "stock_days_actual": days_actual,
                "scope": sheet,
            })
    pct_values = [r[k] for r in records for k in ("shrink_budget_pct", "shrink_actual_pct", "shrink_variance_pct") if r.get(k) is not None]
    raw_numeric = [rows[i][j] for i in range(metric_row + 1, len(rows)) for j in range(len(rows[i]))
                   if isinstance(_cell(rows[i], j), float) and not isinstance(_cell(rows[i], j), Pct)]
    if pct_values and raw_numeric and max(abs(v) for v in pct_values) <= 1.0:
        for r in records:
            for k in ("shrink_budget_pct", "shrink_actual_pct", "shrink_variance_pct"):
                if r.get(k) is not None:
                    r[k] = r[k] * 100
        warnings.append("Shrinkage percentages were stored as fractions in the file and were converted to %.")
    if "out of stock" in sheet.casefold():
        warnings.append("Tab is named 'Out of Stock' but holds shrinkage and stock days; OOS % comes from the sales report.")
    return PackSheet("po_shrink", sheet, records, warnings, f"{len(stores)} stores · {len({r['section'] for r in records})} sections")


def parse_role(rows: list[list[Any]], sheet: str) -> PackSheet:
    header_index = _find_header(rows, "Employee Number", "on this role")
    header = rows[header_index]
    c = lambda *n: _column(header, *n)  # noqa: E731
    cols = {"id": c("Employee Number"), "store": c("Business Unit"), "dept": c("Department"), "section": c("Section"),
            "name": c("Employee Name"), "designation": c("Designation"), "since": c("on this role"), "remarks": c("Remarks")}
    records, warnings = [], []
    no_id = []
    for row in rows[header_index + 1:]:
        name = canon.repair_name(_cell(row, cols["name"]))
        if not name:
            continue
        emp = canon.employee_id(_cell(row, cols["id"]))
        if not emp:
            no_id.append(name)
        remarks = _t(_cell(row, cols["remarks"]))
        section = canon.section(_cell(row, cols["section"]))
        led = canon.sections_in(remarks) or ([section] if section in canon.DEPARTMENT_OF else [])
        since = parse_date(_cell(row, cols["since"]))
        records.append({
            "employee_id": emp, "employee_name": name,
            "store": canon.store(_cell(row, cols["store"])),
            "department": _t(_cell(row, cols["dept"])), "section": section,
            "designation": _t(_cell(row, cols["designation"])) if not str(_cell(row, cols["designation"]) or "").startswith("#") else "",
            "working_as": canon.working_role(remarks, _cell(row, cols["designation"])),
            "role_since": since.isoformat() if since else None,
            "remarks": remarks, "sections_led": led,
            "acting": bool(re.search(r"\bDOA\b|acting", remarks, re.IGNORECASE)),
            "designation_mismatch": bool(re.search(r"but in (the )?system", remarks, re.IGNORECASE)),
        })
    if no_id:
        warnings.append(f"{len(no_id)} leader(s) have no employee number yet: {', '.join(no_id)}.")
    return PackSheet("po_role", sheet, records, warnings, f"{len(records)} leaders")


def parse_da(rows: list[list[Any]], sheet: str) -> PackSheet:
    header_index = _find_header(rows, "Row Labels")
    label_col = _column(rows[header_index], "Row Labels") or 0
    records, warnings = [], []
    current_section = ""
    known = {name for name, *_ in canon._STORES}
    merged_sections: set[str] = set()
    for row in rows[header_index + 1:]:
        label = _t(_cell(row, label_col))
        count = num(_cell(row, label_col + 1))
        if not label or label.casefold().startswith("grand total"):
            continue
        as_store = canon.store(label)
        if as_store.replace(" · Myli", "") in known:
            if current_section and count:
                records.append({"section": canon.section(current_section), "section_raw": current_section,
                                "store": as_store, "store_raw": label, "count": int(count)})
            continue
        if current_section and _t(current_section).casefold() == label.casefold():
            merged_sections.add(label)
        current_section = label
    raw_sections = [r["section_raw"] for r in records]
    dupes = {s for s in raw_sections if raw_sections.count(s) and any(o != s and o.casefold() == s.casefold() for o in raw_sections)}
    if dupes or merged_sections:
        warnings.append("Some section labels appeared twice because of trailing spaces; they were merged.")
    warnings.append("Source is a pivot table: counts by section and store only, with no names or dates.")
    total = sum(r["count"] for r in records)
    return PackSheet("po_da", sheet, records, warnings, f"{total} cases · {len({r['store'] for r in records})} stores")


def parse_turnover(rows: list[list[Any]], sheet: str) -> PackSheet:
    header_index = _find_header(rows, "End of Service Type")
    header = rows[header_index]
    c = lambda *n: _column(header, *n)  # noqa: E731
    cols = {"id": c("Emp. No.", "Employee Number"), "name": c("Emp. Name", "Employee Name"),
            "date": c("Requested Date"), "store": c("Business Unit"), "dept": c("Dept."), "section": c("Sect."),
            "type": c("End of Service Type"), "reason": c("Reason"), "salary": c("Salary Period")}
    records = []
    odd_periods = 0
    for row in rows[header_index + 1:]:
        emp = canon.employee_id(_cell(row, cols["id"]))
        name = _t(_cell(row, cols["name"]))
        if not emp and not name:
            continue
        exit_date = parse_date(_cell(row, cols["date"]))
        salary = _t(_cell(row, cols["salary"]))
        salary = salary[:-2] if salary.endswith(".0") else salary
        reason = re.sub(r"[ÂÃ ]+", " ", _t(_cell(row, cols["reason"])))
        reason = re.sub(r"\s+", " ", reason).strip()
        reason = re.sub(r"^(Voluntary|Involuntary)\s+", "", reason)
        section_raw = _t(_cell(row, cols["section"]))
        dept_raw = _t(_cell(row, cols["dept"]))
        section = canon.section(section_raw)
        if section in {"", section_raw}:
            section = canon.section(dept_raw)
        if exit_date and re.fullmatch(r"\d{6}", salary) and salary != f"{exit_date:%Y%m}":
            odd_periods += 1
        exit_type = _t(_cell(row, cols["type"]))
        records.append({
            "employee_id": emp, "employee_name": canon.repair_name(name).rstrip(" ."),
            "date": exit_date.isoformat() if exit_date else None,
            "store": canon.store(_cell(row, cols["store"])), "department": dept_raw,
            "section": section, "section_raw": section_raw,
            "exit_type": exit_type, "reason": reason, "salary_period": salary,
            "voluntary": exit_type.casefold() == "resignation",
        })
    warnings = []
    if odd_periods:
        warnings.append(f"{odd_periods} leaver(s) have a salary period different from the requested-date month; counts use the requested date.")
    dated = sorted(r["date"] for r in records if r["date"])
    span = f"{dated[0][:7]} – {dated[-1][:7]}" if dated else "undated"
    return PackSheet("po_turnover", sheet, records, warnings, f"{len(records)} leavers · {span}")


def parse_master(rows: list[list[Any]], sheet: str) -> PackSheet:
    header_index = _find_header(rows, "Employee Number", "Joining Date")
    header = rows[header_index]
    c = lambda *n: _column(header, *n)  # noqa: E731
    cols = {"id": c("Employee Number"), "store": c("Business Unit"), "dept": c("Department"), "section": c("Section"),
            "name": c("Employee Name"), "designation": c("Designation"), "grade": c("Grade"), "gender": c("Gender"),
            "manager": c("Reporting Manager"), "joining": c("Joining Date")}
    records = []
    for row in rows[header_index + 1:]:
        emp = canon.employee_id(_cell(row, cols["id"]))
        if not emp:
            continue
        joining = parse_date(_cell(row, cols["joining"]))
        grade = _t(_cell(row, cols["grade"]))
        grade = grade.lstrip("0") or grade
        records.append({
            "employee_id": emp, "employee_name": canon.repair_name(_cell(row, cols["name"])),
            "store": canon.store(_cell(row, cols["store"])), "store_raw": _t(_cell(row, cols["store"])),
            "department": _t(_cell(row, cols["dept"])),
            "section": canon.section(_cell(row, cols["section"])), "section_raw": _t(_cell(row, cols["section"])),
            "designation": _t(_cell(row, cols["designation"])), "grade": grade,
            "gender": _t(_cell(row, cols["gender"])),
            "manager_name": canon.repair_name(_cell(row, cols["manager"])),
            "joining_date": joining.isoformat() if joining else None,
        })
    by_name: dict[str, list[str]] = {}
    for record in records:
        by_name.setdefault(canon.name_key(record["employee_name"]), []).append(record["employee_id"])
    resolved = 0
    for record in records:
        matches = by_name.get(canon.name_key(record["manager_name"]), [])
        record["manager_id"] = matches[0] if len(matches) == 1 else ""
        resolved += bool(record["manager_id"])
    warnings = ["Birth date, age, e-mail, mobile number and marital status were not imported (not needed for analysis)."]
    with_manager = sum(1 for r in records if r["manager_name"])
    if with_manager and resolved < with_manager:
        warnings.append(f"Manager linked by ID for {resolved} of {with_manager} people; the rest report to someone outside this sheet.")
    return PackSheet("po_master", sheet, records, warnings, f"{len(records)} employees · {len({r['store'] for r in records})} stores")


# ----------------------------------------------------------- detection --

_DETECTORS: list[tuple[str, Callable[[str], bool], Callable[[list[list[Any]], str], PackSheet]]] = [
    ("po_sales", lambda s: "section code name" in s or "periodic store performance" in s, parse_sales),
    ("po_productivity", lambda s: "store productivity" in s and "target productivity" in s, parse_productivity),
    ("po_rating", lambda s: "subject employee id" in s or "form template name" in s, parse_rating),
    ("po_turnover", lambda s: "end of service type" in s, parse_turnover),
    ("po_role", lambda s: "on this role" in s, parse_role),
    ("po_master", lambda s: "joining date" in s and "employee number" in s, parse_master),
    ("po_nps", lambda s: "product influence" in s and "nps" in s, parse_nps),
    ("po_shrink", lambda s: ("shrink" in s and "stock days" in s), parse_shrink),
    ("po_da", lambda s: "row labels" in s and "count of" in s, parse_da),
    ("po_shift", lambda s: "employee number" in s and "designation" in s and bool(re.search(r"\b\d{1,2} (jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\w* \d{2}", s) or re.search(r"20\d\d \d\d \d\d", s)), parse_roster),
]


def _signature(rows: list[list[Any]]) -> str:
    parts = []
    for row in rows[:45]:
        for value in row:
            if isinstance(value, (datetime, date)):
                parts.append(value.strftime("%d %b %y").lower())
            elif value is not None:
                parts.append(_norm(value))
    return " | ".join(parts)


def detect(rows: list[list[Any]]) -> str | None:
    signature = _signature(rows)
    for kind, test, _parser in _DETECTORS:
        if test(signature):
            return kind
    return None


def parse_sheet(name: str, rows: list[list[Any]], kind: str | None = None) -> PackSheet | None:
    kind = kind or detect(rows)
    for candidate, _test, parser in _DETECTORS:
        if candidate == kind:
            return parser(rows, name)
    return None


def scan_pack(path: str | Path) -> PackPlan:
    source = Path(path)
    checksum = hashlib.sha256(source.read_bytes()).hexdigest()
    sheets, skipped = [], []
    for name, rows in read_sheets(source):
        try:
            parsed = parse_sheet(name, rows)
        except Exception as exc:  # a malformed tab must not block the others
            skipped.append(f"{name}: could not be read ({exc})")
            continue
        if parsed is None or not parsed.records:
            skipped.append(f"{name}: not a recognised P&O tab")
            continue
        sheets.append(parsed)
    hint = None
    for sheet in sheets:
        if sheet.kind == "po_sales":
            ends = [r["period_to"] for r in sheet.records if r.get("period_to")]
            if ends:
                hint = ends[0][:7]
    return PackPlan(str(source), checksum, sheets, skipped, hint)


def roster_mirror(master: PackSheet, period: str) -> list[dict[str, Any]]:
    """Basic Sheet rows in the existing roster shape, so Staff & hierarchy uses them."""
    as_of = f"{period}-01" if len(period) == 7 else period
    return [
        {
            "store": r["store"], "department": r["department"], "section": r["section_raw"] or r["section"],
            "employee_id": r["employee_id"], "employee_name": r["employee_name"],
            "designation": r["designation"], "grade": r["grade"], "gender": r["gender"],
            "manager_id": r["manager_id"], "manager_name": r["manager_name"], "date": as_of,
        }
        for r in master.records
    ]


def import_pack(store: Any, plan: PackPlan, period: str, selected: list[str] | None = None) -> list[str]:
    """Save every selected tab. Re-importing a month replaces that month's batch."""
    messages = []
    for sheet in plan.sheets:
        if selected is not None and sheet.kind not in selected:
            continue
        checksum = hashlib.sha256(f"{plan.checksum}:{sheet.sheet}:{sheet.kind}".encode()).hexdigest()
        try:
            store.import_records(sheet.kind, period, sheet.records, Path(plan.path).name, checksum, sheet.sheet, replace=True)
            messages.append(f"✓ {sheet.label}: {sheet.summary}")
        except Exception as exc:
            messages.append(f"• {sheet.label}: {exc}")
        if sheet.kind == "po_master":
            try:
                store.import_records("roster", f"{period}-01" if len(period) == 7 else period,
                                     roster_mirror(sheet, period), Path(plan.path).name, checksum, sheet.sheet, replace=True)
                messages.append("✓ Staff & hierarchy updated from the Basic Sheet")
            except Exception as exc:
                messages.append(f"• Staff & hierarchy not updated: {exc}")
    return messages
