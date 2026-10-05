"""Corporate exports: one report model, rendered to PDF (A4), PowerPoint (16:9, editable charts) or Excel.

Packs: director (Monthly HR Director Pack), store, dept (department / section comparison), scorecard (individual or all
managers in scope), attendance (exceptions), people (ranking table). Every export prints its scope, month and sources.
"""

from __future__ import annotations

import re
from datetime import datetime
from typing import Any

from . import analytics, engine
from .db import Database
from .org import DEPT_NAMES, ROLE_LABEL
from .util import fmt_minutes, month_label

BROWN, BROWN2, GOLD, GOLD_SOFT, INK, MUTED, LINE, BG = "#2A1A11", "#94481A", "#B8860B", "#F6ECD4", "#1F1915", "#766C63", "#E7E2DC", "#FAF8F5"
GOOD, WARN, BAD = "#2E7D4F", "#B7791F", "#B3261E"
BAND_COLORS = {"excellent": "#2E7D4F", "good": "#5B9A63", "ok": "#B8860B", "warn": "#D9822B", "bad": "#B3261E", "none": "#9A8F84"}
PACKS = {"director": "Monthly HR Director Pack", "store": "Store Report", "dept": "Department Comparison",
         "scorecard": "Individual Scorecard", "attendance": "Attendance Exceptions", "people": "People Ranking"}


def _f(v, fmt: str = "") -> str:
    if v is None or v == "":
        return "—"
    if fmt == "pct":
        return f"{v:.1f}%"
    if fmt == "pct0":
        return f"{v:.0f}%"
    if fmt == "score":
        return f"{v:.1f}"
    if fmt == "money":
        return f"{v:,.0f}"
    if fmt == "int":
        return f"{v:,.0f}"
    if fmt == "signed":
        return f"{v:+.1f}"
    return str(v)


def _name(p: dict, hide: bool) -> str:
    return f"Employee {p['emp']}" if hide else p["name"]


# ------------------------------------------------------------------------------------------------ packs
def build_pack(db: Database, pack: str, month: str | None, scope: dict, emp=None, store_id=None, hide=False) -> dict:
    if pack not in PACKS:
        raise ValueError(f"Unknown report: {pack}")
    mm = engine.model(db, month)
    rep: dict[str, Any] = {"pack": pack, "title": PACKS[pack], "month": mm.month, "month_label": month_label(mm.month),
                           "scope": analytics.scope_label(mm, scope), "generated": datetime.now().strftime("%d %b %Y %H:%M"),
                           "hide": hide, "sections": [], "as_of": mm.as_of}
    rep["sources"] = [f"{r['kind'].title()}: {r['file_name']} ({r['as_of'] or r['month'] or ''})" for r in db.q(
        "SELECT kind, file_name, as_of, month FROM imports WHERE active=1 AND (month=? OR kind='roster') AND kind!='leave' "
        "ORDER BY kind, id DESC", (mm.month,))][:12]
    globals()[f"_pack_{pack}"](db, mm, rep, scope, emp, store_id, hide)
    return rep


def _kpis_from_cards(cards) -> dict:
    items = []
    for c in cards:
        fmt = c["fmt"]
        val = _f(c["value"], "score" if fmt == "score" else fmt)
        d = c.get("delta")
        sub = c.get("sub") or ""
        if d is not None:
            ds = f"{d:+,.0f}" if fmt == "int" else f"{d:+.1f}" + (" pts" if fmt.startswith("pct") else "")
            sub = f"{ds} vs last month" + (f" · {sub}" if sub else "")
        items.append({"label": c["label"], "value": val, "sub": sub})
    return {"type": "kpis", "items": items}


def _pack_director(db, mm, rep, scope, emp, store_id, hide):
    d = analytics.dashboard(db, mm.month, scope, None, hide)
    rep["subtitle"] = f"{d['scope_label']} · {d['month_label']}" + (" · provisional" if d["provisional"] else "")
    rep["sections"].append(_kpis_from_cards(d["cards"]))
    if d["trend"]:
        labels = [t["label"] for t in d["trend"]]
        rep["sections"].append({"type": "chart", "kind": "line", "title": "Average score trend (out of 10)", "labels": labels,
                                "series": [{"name": "Average score", "values": [t["score"] for t in d["trend"]]}]})
        rep["sections"].append({"type": "chart", "kind": "line", "title": "Real absence trend (%)", "labels": labels,
                                "series": [{"name": "Real absence %", "values": [t["absence"] for t in d["trend"]]}]})
        if any(t["sales"] is not None for t in d["trend"]):
            rep["sections"].append({"type": "chart", "kind": "line", "title": "Sales vs budget trend (%)", "labels": labels,
                                    "series": [{"name": "% of budget", "values": [t["sales"] for t in d["trend"]]}]})
    rep["sections"].append({"type": "chart", "kind": "bar", "title": "Score distribution",
                            "labels": [x["label"] for x in d["distribution"]],
                            "series": [{"name": "People", "values": [x["count"] for x in d["distribution"]]}],
                            "colors": [BAND_COLORS[x["key"]] for x in d["distribution"]]})
    league = d["league"]
    if league:
        rep["sections"].append({"type": "chart", "kind": "hbar", "title": "Store league · average score",
                                "labels": [r["store"] for r in league[:15]],
                                "series": [{"name": "Average score", "values": [r["score"] for r in league[:15]]}]})
        rep["sections"].append({"type": "table", "title": "Store league", "columns": [
            ("Store", "store", ""), ("Format", "format", ""), ("People", "people", "int"), ("Avg score", "score", "score"),
            ("GM", "gm", ""), ("GM score", "gm_score", "score"), ("Sales vs budget", "sales_pct", "pct0"),
            ("Growth", "growth", "signed"), ("Productivity", "prod_pct", "pct0"), ("Real absence", "absence_pct", "pct")],
            "rows": [{**r, "gm": (f"Employee {r['gm_emp']}" if hide and r["gm_emp"] else r["gm"])} for r in league]})
    cols = [("Name", "name", ""), ("Role", "role_label", ""), ("Store", "store", ""), ("Score", "score", "score"), ("Band", "band", "")]
    rep["sections"].append({"type": "table", "title": "Top 10 performers", "columns": cols, "rows": d["top"]})
    rep["sections"].append({"type": "table", "title": "Needs attention · bottom 10", "columns": cols, "rows": d["bottom"]})
    a = d["attendance"]
    rep["sections"].append({"type": "kpis", "title": "Attendance health", "items": [
        {"label": "Real absence", "value": _f(a["absence_pct"], "pct"), "sub": f"System shows {_f(a['system_absence_pct'], 'pct')}"},
        {"label": "Full 9-hour days", "value": _f(a["compliance_pct"], "pct0"), "sub": f"Average day {fmt_minutes(a['avg_minutes'])}"},
        {"label": "Punches to fix", "value": _f(a["fixes"], "int"), "sub": "Missing IN / OUT"},
        {"label": "Long absences", "value": _f(a["long_absence"], "int"), "sub": "7+ days in a row"}]})
    if d["alerts"]:
        rep["sections"].append({"type": "text", "title": "Things to know", "lines": [x["text"] for x in d["alerts"]]})


def _pack_store(db, mm, rep, scope, emp, store_id, hide):
    sid = store_id or scope.get("store_id")
    if not sid:
        raise ValueError("Choose a store first (Store filter), then export the store report.")
    s = analytics.store_detail(db, mm.month, int(sid), hide)
    rep["title"] = f"Store Report · {s['store']}"
    rep["subtitle"] = f"{s['format_label']} · {s['city']} · {s['month_label']}"
    t = s["total"] or {}
    gm = s["gm"]
    rep["sections"].append({"type": "kpis", "items": [
        {"label": "Store score", "value": _f(s["score"], "score"), "sub": f"{s['people']} people"},
        {"label": "Store manager", "value": _f(gm and gm["score"], "score"), "sub": _name(gm, hide) if gm else "Not in roster"},
        {"label": "Sales vs budget", "value": _f(t.get("sales_pct"), "pct0"), "sub": f"{_f(t.get('growth'), 'signed')}% vs LY"},
        {"label": "Real absence", "value": _f(s["attendance"]["absence_pct"], "pct"), "sub": f"{s['attendance']['fixes']} punches to fix"}]})
    if s["gm_components"]:
        rep["sections"].append({"type": "table", "title": "Store manager scorecard", "columns": [
            ("Part", "label", ""), ("Weight", "weight", "int"), ("Result", "display", ""), ("Score /10", "score", "score")],
            "rows": s["gm_components"]})
    if s["departments"]:
        rep["sections"].append({"type": "table", "title": "Departments", "columns": [
            ("Code", "code", ""), ("Department", "name", ""), ("Head", "mgr", ""), ("Sales", "actual", "money"),
            ("vs budget", "sales_pct", "pct0"), ("Growth", "growth", "signed"), ("Margin", "margin", "pct"),
            ("Waste", "waste", "pct"), ("OOS", "oos", "pct"), ("Productivity", "prod_pct", "pct0")],
            "rows": [{**r, "mgr": ", ".join(r["managers"]) or "—"} for r in s["departments"]]})
    if s["sections"]:
        rep["sections"].append({"type": "table", "title": "Sections and their managers", "columns": [
            ("Code", "code", ""), ("Section", "name", ""), ("Manager", "mgr", ""), ("Sales", "actual", "money"),
            ("vs budget", "sales_pct", "pct0"), ("Growth", "growth", "signed"), ("Waste", "waste", "pct"), ("OOS", "oos", "pct"),
            ("Productivity", "prod_pct", "pct0")],
            "rows": [{**r, "mgr": ", ".join(r["managers"]) or "—"} for r in s["sections"]]})
    if s["managers"]:
        rep["sections"].append({"type": "table", "title": "Managers", "columns": [
            ("Name", "name", ""), ("Designation", "designation", ""), ("Role", "role_label", ""), ("Score", "score", "score"),
            ("Band", "band", ""), ("Attendance", "attendance", "score")], "rows": s["managers"]})
    rep["sections"].append({"type": "table", "title": "Attendance by section", "columns": [
        ("Section", "section", ""), ("People", "people", "int"), ("Worked days", "worked", "int"), ("Real absence", "absence_pct", "pct"),
        ("Full 9h days", "compliance_pct", "pct0"), ("Punches to fix", "fixes", "int")], "rows": s["attendance_sections"]})


def _pack_dept(db, mm, rep, scope, emp, store_id, hide):
    dept = scope.get("dept")
    if not dept:
        raise ValueError("Choose a department first (Department filter), then export the comparison.")
    rep["title"] = f"{DEPT_NAMES.get(dept, dept)} · Store Comparison"
    rep["subtitle"] = f"{analytics.scope_label(mm, {k: v for k, v in scope.items() if k != 'dept'})} · {month_label(mm.month)}"
    rows = []
    for sid, st in mm.stores.items():
        if not analytics.store_in_scope(st, scope):
            continue
        r = getattr(mm, "sales", {}).get(sid, {}).get(f"D{dept}")
        pr = getattr(mm, "prod", {}).get(sid, {}).get(f"D{dept}")
        heads = [p for p in mm.people.values() if p.store_id == sid and p.role == "DH" and p.domain.get("dept") == dept]
        if not r and not heads:
            continue
        rows.append({"store": st["name"], "format": st["format"],
                     "head": ", ".join(_name(h.brief(), hide) for h in heads) or "—",
                     "head_score": heads[0].score if heads else None, "actual": r and r["actual"],
                     "sales_pct": r["actual"] / r["budget"] * 100 if r and r.get("budget") else None,
                     "growth": r and r["growth"], "waste": r and r["waste"], "oos": r and r["oos"],
                     "prod_pct": pr["actual"] / pr["target"] * 100 if pr and pr.get("target") else None})
    rows.sort(key=lambda x: -(x["sales_pct"] or 0))
    rep["sections"].append({"type": "chart", "kind": "hbar", "title": "Sales vs budget by store (%)", "labels": [r["store"] for r in rows],
                            "series": [{"name": "% of budget", "values": [r["sales_pct"] for r in rows]}]})
    rep["sections"].append({"type": "table", "title": "Store comparison", "columns": [
        ("Store", "store", ""), ("Format", "format", ""), ("Department head", "head", ""), ("Head score", "head_score", "score"),
        ("Sales", "actual", "money"), ("vs budget", "sales_pct", "pct0"), ("Growth", "growth", "signed"), ("Waste", "waste", "pct"),
        ("OOS", "oos", "pct"), ("Productivity", "prod_pct", "pct0")], "rows": rows})
    ppl = sorted([p for p in mm.people.values() if analytics.in_scope(p, scope) and p.role in ("SM", "DH") and p.score is not None],
                 key=lambda p: -p.score)
    rep["sections"].append({"type": "table", "title": "Managers in this department", "columns": [
        ("Name", "name", ""), ("Designation", "designation", ""), ("Store", "store", ""), ("Score", "score", "score"), ("Band", "band", "")],
        "rows": [analytics._person_row(p, hide) for p in ppl]})


def _scorecard_sections(db, mm, emp, hide) -> list[dict]:
    p = analytics.person(db, mm.month, emp, hide)
    secs: list[dict] = [{"type": "heading", "title": p["name"],
                         "sub": f"{p['designation']} · {p['store']} · {p['role_label']} · Employee {p['emp']}"}]
    a = p["attendance_detail"] or {}
    secs.append({"type": "kpis", "items": [
        {"label": "Score", "value": _f(p["score"], "score") if p["score"] is not None else
         {"new_joiner": "New joiner", "not_enough_data": "Not enough data"}.get(p["status"], "—"),
         "sub": p["band"] + (" · provisional" if p["provisional"] else "")},
        {"label": "Measured on", "value": f"{p['coverage']}%", "sub": p["domain"]},
        {"label": "Presence", "value": _f(a.get("presence"), "pct0"), "sub": f"{a.get('worked', 0)} days worked · {a.get('absent', 0)} absent"},
        {"label": "Full 9h days", "value": _f(a.get("compliance"), "pct0"), "sub": f"Average {a.get('avg_hours') or '—'}"}]})
    secs.append({"type": "table", "title": "How the score is made", "columns": [
        ("Part", "label", ""), ("Weight", "weight", "int"), ("Result", "display", ""), ("Score /10", "score", "score"), ("Note", "detail", "")],
        "rows": [{**c, "display": c["display"] or "No data (left out)"} for c in p["components"]]})
    if p["context"]:
        secs.append({"type": "text", "title": "Context (not scored)", "lines": [f"{c['key'].title()}: {c['display']}" for c in p["context"]]})
    if p["history"] and len(p["history"]) > 1:
        secs.append({"type": "chart", "kind": "line", "title": "Score history", "labels": [h["label"] for h in p["history"]],
                     "series": [{"name": "Score", "values": [h["score"] for h in p["history"]]},
                                {"name": "Attendance score", "values": [h["attendance"] for h in p["history"]]}]})
    if a.get("calendar"):
        kinds = {"worked": "Worked", "absent": "Absent", "off": "Weekly off", "leave": "Leave", "holiday": "Holiday"}
        st = {"missing_in": "Missing IN", "missing_out": "Missing OUT", "punched_absent": "Punched, marked absent", "manual": "Manual hours"}
        secs.append({"type": "table", "title": "Attendance days", "small": True, "columns": [
            ("Date", "day", ""), ("Day", "kind", ""), ("IN", "t_in", ""), ("OUT", "t_out", ""), ("Hours", "hours", ""), ("Note", "note", "")],
            "rows": [{**c, "kind": kinds.get(c["kind"], c["kind"]),
                      "note": st.get(c["status"], "") or c.get("leave") or c.get("holiday") or ""} for c in a["calendar"]]})
    return secs


def _pack_scorecard(db, mm, rep, scope, emp, store_id, hide):
    if emp:
        rep["sections"] += _scorecard_sections(db, mm, str(emp), hide)
        p = mm.people.get(str(emp))
        rep["title"] = "Individual Scorecard"
        rep["subtitle"] = f"{(f'Employee {emp}' if hide else p.name) if p else emp} · {month_label(mm.month)}"
        return
    ppl = sorted([p for p in mm.people.values() if analytics.in_scope(p, scope) and p.role != "STAFF"],
                 key=lambda p: (p.store, -(p.score or -1)))[:200]
    if not ppl:
        raise ValueError("No managers in this scope.")
    rep["title"] = "Manager Scorecards"
    rep["subtitle"] = f"{rep['scope']} · {len(ppl)} managers · {month_label(mm.month)}"
    for i, p in enumerate(ppl):
        if i:
            rep["sections"].append({"type": "pagebreak"})
        rep["sections"] += [s for s in _scorecard_sections(db, mm, p.emp, hide) if s.get("title") != "Attendance days"]


def _pack_attendance(db, mm, rep, scope, emp, store_id, hide):
    a = analytics.attendance_view(db, mm.month, scope, hide)
    t = a["totals"]
    rep["subtitle"] = f"{rep['scope']} · {a['month_label']} (up to {a['as_of'] or '—'})"
    rep["sections"].append({"type": "kpis", "items": [
        {"label": "Real absence", "value": _f(t["absence_pct"], "pct"), "sub": f"System shows {_f(t['system_absence_pct'], 'pct')}"},
        {"label": "Full 9-hour days", "value": _f(t["compliance_pct"], "pct0"), "sub": f"Average day {fmt_minutes(t['avg_minutes'])}"},
        {"label": "Punches to fix", "value": _f(t["fixes"], "int"), "sub": f"{t['people']} people in scope"},
        {"label": "Long absences", "value": _f(t["long_absence"], "int"), "sub": "7+ days in a row"}]})
    if a["daily"]:
        rep["sections"].append({"type": "chart", "kind": "line", "title": "People at work per day",
                                "labels": [d["label"] for d in a["daily"]],
                                "series": [{"name": "Worked", "values": [d["worked"] for d in a["daily"]]},
                                           {"name": "Absent", "values": [d["absent"] for d in a["daily"]]}]})
    rep["sections"].append({"type": "table", "title": "By store", "columns": [
        ("Store", "store", ""), ("People", "people", "int"), ("Real absence", "absence_pct", "pct"), ("System absence", "system_absence_pct", "pct"),
        ("Full 9h days", "compliance_pct", "pct0"), ("Punches to fix", "fixes", "int"), ("Long absences", "long_absence", "int")],
        "rows": a["by_store"]})
    if a["long_absence"]:
        rep["sections"].append({"type": "table", "title": "Long absences (possible leavers / long leave)", "columns": [
            ("Name", "name", ""), ("Store", "store", ""), ("Section", "section", ""), ("Days in a row", "longest", "int"),
            ("Absent", "absent", "int"), ("Worked", "worked", "int")], "rows": a["long_absence"]})
    if a["fixes"]:
        kinds = {"missing_in": "Missing IN", "missing_out": "Missing OUT", "punched_absent": "Punched, marked absent"}
        rep["sections"].append({"type": "table", "title": "Punches to fix", "small": True, "columns": [
            ("Store", "store", ""), ("Name", "name", ""), ("Date", "day", ""), ("Problem", "kind", ""), ("IN", "t_in", ""), ("OUT", "t_out", "")],
            "rows": [{**f, "kind": kinds[f["kind"]]} for f in a["fixes"][:1500]]})
    rep["sections"].append({"type": "table", "title": "Lowest 9-hour compliance", "columns": [
        ("Name", "name", ""), ("Store", "store", ""), ("Full 9h days", "compliance", "pct0"), ("Average day", "avg_hours", ""),
        ("Worked", "worked", "int")], "rows": a["low_compliance"]})


def _pack_people(db, mm, rep, scope, emp, store_id, hide):
    data = analytics.people_list(db, mm.month, scope, limit=100000, hide=hide)
    rep["subtitle"] = f"{rep['scope']} · {data['total']} people · {month_label(mm.month)}"
    rep["sections"].append({"type": "table", "title": "People ranking", "columns": [
        ("Employee", "emp", ""), ("Name", "name", ""), ("Designation", "designation", ""), ("Role", "role_label", ""),
        ("Store", "store", ""), ("Department", "dept", ""), ("Score", "score", "score"), ("Band", "band", ""),
        ("Change", "delta", "signed"), ("Attendance", "attendance", "score"), ("Presence", "presence", "pct0"),
        ("Absent days", "absent", "int"), ("Punches to fix", "fixes", "int")], "rows": data["rows"]})


def file_name(rep: dict, fmt: str) -> str:
    base = re.sub(r"[^\w\- ]+", "", f"PNO {rep['title']} {rep['month_label']}").strip()
    if fmt not in ("pdf", "pptx", "xlsx"):
        raise ValueError("Choose PDF, PowerPoint or Excel.")
    return f"{base}.{fmt}"


def render(rep: dict, fmt: str, path: str) -> None:
    {"pdf": render_pdf, "pptx": render_pptx, "xlsx": render_xlsx}[fmt](rep, path)


# ------------------------------------------------------------------------------------------------ Excel
def render_xlsx(rep: dict, path: str) -> None:
    import xlsxwriter
    wb = xlsxwriter.Workbook(path)
    title = wb.add_format({"bold": True, "font_size": 18, "font_color": BROWN})
    sub = wb.add_format({"font_color": MUTED})
    head = wb.add_format({"bold": True, "bg_color": BROWN, "font_color": "white", "border": 1, "border_color": LINE, "text_wrap": True, "valign": "vcenter"})
    cell = wb.add_format({"border": 1, "border_color": LINE})
    fmts = {"pct": wb.add_format({"num_format": "0.0\"%\"", "border": 1, "border_color": LINE}),
            "pct0": wb.add_format({"num_format": "0\"%\"", "border": 1, "border_color": LINE}),
            "score": wb.add_format({"num_format": "0.0", "border": 1, "border_color": LINE}),
            "money": wb.add_format({"num_format": "#,##0", "border": 1, "border_color": LINE}),
            "int": wb.add_format({"num_format": "#,##0", "border": 1, "border_color": LINE}),
            "signed": wb.add_format({"num_format": "+0.0;-0.0;0.0", "border": 1, "border_color": LINE})}
    kpi_l = wb.add_format({"bold": True, "font_color": MUTED, "font_size": 9})
    kpi_v = wb.add_format({"bold": True, "font_size": 16, "font_color": BROWN})
    ws = wb.add_worksheet("Summary")
    ws.set_column(0, 0, 28)
    ws.set_column(1, 6, 22)
    ws.write(0, 0, f"PNO · {rep['title']}", title)
    ws.write(1, 0, rep.get("subtitle") or f"{rep['scope']} · {rep['month_label']}", sub)
    ws.write(2, 0, f"Generated {rep['generated']} · Confidential – HR" + (" · names hidden" if rep["hide"] else ""), sub)
    r = 4
    used = {"Summary"}
    for s in rep["sections"]:
        if s["type"] == "kpis":
            if s.get("title"):
                ws.write(r, 0, s["title"], wb.add_format({"bold": True, "font_color": BROWN}))
                r += 1
            for i, it in enumerate(s["items"]):
                ws.write(r, i, it["label"], kpi_l)
                ws.write(r + 1, i, it["value"], kpi_v)
                ws.write(r + 2, i, it.get("sub", ""), sub)
            r += 4
        elif s["type"] in ("text", "heading"):
            ws.write(r, 0, s["title"], wb.add_format({"bold": True, "font_color": BROWN}))
            r += 1
            for line in s.get("lines", [s.get("sub", "")]):
                ws.write(r, 0, line)
                r += 1
            r += 1
        elif s["type"] in ("table", "chart"):
            name = re.sub(r"[\[\]:*?/\\]", "", s["title"])[:28] or "Table"
            base, k = name, 2
            while name in used:
                name = f"{base[:25]} {k}"
                k += 1
            used.add(name)
            sh = wb.add_worksheet(name)
            sh.write(0, 0, s["title"], title)
            sh.write(1, 0, f"{rep['scope']} · {rep['month_label']}", sub)
            if s["type"] == "table":
                cols = s["columns"]
                for j, (lbl, key, fm) in enumerate(cols):
                    sh.write(3, j, lbl, head)
                    sh.set_column(j, j, max(10, min(40, len(lbl) + 6)) if fm else 26)
                for i, row in enumerate(s["rows"], 4):
                    for j, (lbl, key, fm) in enumerate(cols):
                        v = row.get(key)
                        if v is None or v == "":
                            sh.write_blank(i, j, None, cell)
                        elif fm and isinstance(v, (int, float)):
                            sh.write_number(i, j, v, fmts[fm])
                        else:
                            sh.write_string(i, j, str(v), cell)
                sh.autofilter(3, 0, 3 + len(s["rows"]), len(cols) - 1)
                sh.freeze_panes(4, 0)
            else:
                sh.write(3, 0, "", head)
                for j, se in enumerate(s["series"], 1):
                    sh.write(3, j, se["name"], head)
                for i, lab in enumerate(s["labels"], 4):
                    sh.write(i, 0, lab, cell)
                    for j, se in enumerate(s["series"], 1):
                        v = se["values"][i - 4]
                        if v is None:
                            sh.write_blank(i, j, None, cell)
                        else:
                            sh.write_number(i, j, v, fmts["score"])
                sh.set_column(0, 0, 26)
                ctype = {"line": "line", "bar": "column", "hbar": "bar"}[s["kind"]]
                ch = wb.add_chart({"type": ctype})
                palette = [BROWN2, "#3B6EA5", GOLD, "#6C8E3A"]
                n = len(s["labels"])
                for j, se in enumerate(s["series"], 1):
                    spec = {"name": se["name"], "categories": [name, 4, 0, 3 + n, 0], "values": [name, 4, j, 3 + n, j]}
                    if ctype == "line":
                        spec["line"] = {"color": palette[(j - 1) % 4], "width": 2.5}
                        spec["marker"] = {"type": "circle", "size": 6, "fill": {"color": palette[(j - 1) % 4]}, "border": {"color": "white"}}
                    else:
                        spec["fill"] = {"color": palette[(j - 1) % 4]}
                        if s.get("colors") and len(s["series"]) == 1:
                            spec["points"] = [{"fill": {"color": c}} for c in s["colors"]]
                    ch.add_series(spec)
                ch.set_title({"name": s["title"], "name_font": {"color": BROWN, "size": 12}})
                ch.set_legend({"position": "bottom"})
                ch.set_size({"width": 760, "height": 380})
                sh.insert_chart(3, len(s["series"]) + 2, ch)
    if rep.get("sources"):
        r += 1
        ws.write(r, 0, "Sources", wb.add_format({"bold": True, "font_color": BROWN}))
        for line in rep["sources"]:
            r += 1
            ws.write(r, 0, line, sub)
    wb.close()


# ------------------------------------------------------------------------------------------------ PDF
def render_pdf(rep: dict, path: str) -> None:
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    C = colors.HexColor
    page = landscape(A4)
    W = page[0] - 30 * mm
    st_title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=20, textColor=C(BROWN), leading=24)
    st_sub = ParagraphStyle("s", fontName="Helvetica", fontSize=10, textColor=C(MUTED), leading=13)
    st_h = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=12.5, textColor=C(BROWN), leading=16, spaceBefore=8, spaceAfter=4)
    st_cell = ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=10, textColor=C(INK))
    st_cell_s = ParagraphStyle("cs", fontName="Helvetica", fontSize=7, leading=8.5, textColor=C(INK))
    st_txt = ParagraphStyle("x", fontName="Helvetica", fontSize=9.5, leading=13, textColor=C(INK))

    def esc(s):
        return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story: list = [Paragraph(f"PNO · {esc(rep['title'])}", st_title),
                   Paragraph(esc(rep.get("subtitle") or f"{rep['scope']} · {rep['month_label']}"), st_sub), Spacer(1, 6 * mm)]

    def kpis(items):
        n = len(items)
        cw = W / n
        cells = [[Paragraph(f"<font size=8 color='{MUTED}'><b>{esc(it['label']).upper()}</b></font><br/>"
                            f"<font size=17 color='{BROWN}'><b>{esc(it['value'])}</b></font><br/>"
                            f"<font size=7.5 color='{MUTED}'>{esc(it.get('sub', ''))}</font>",
                            ParagraphStyle("k", leading=20)) for it in items]]
        t = Table(cells, colWidths=[cw - 2 * mm] * n)
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, C(LINE)), ("INNERGRID", (0, 0), (-1, -1), 0.6, C(LINE)),
                               ("BACKGROUND", (0, 0), (-1, -1), colors.white), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 6),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 8), ("LINEABOVE", (0, 0), (-1, 0), 2.5, C(GOLD))]))
        return t

    def table(s):
        cols = s["columns"]
        small = s.get("small") or len(cols) > 9
        sty = st_cell_s if small else st_cell
        data = [[Paragraph(f"<b>{esc(c[0])}</b>", ParagraphStyle("hd", parent=sty, textColor=colors.white)) for c in cols]]
        for row in s["rows"]:
            data.append([Paragraph(esc(_f(row.get(k), fm) if fm else (row.get(k) if row.get(k) not in (None, "") else "—")), sty)
                         for _, k, fm in cols])
        weights = [1.0 if fm else 2.0 for _, _, fm in cols]
        tot = sum(weights)
        t = Table(data, colWidths=[W * w / tot for w in weights], repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), C(BROWN)), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, C(BG)]),
                               ("LINEBELOW", (0, 0), (-1, -1), 0.4, C(LINE)), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                               ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        return t

    def chart(s):
        h = 62 * mm if s["kind"] != "hbar" else max(50 * mm, 7 * mm * len(s["labels"]) + 20 * mm)
        d = Drawing(W, h)
        palette = [C(BROWN2), C("#3B6EA5"), C(GOLD), C("#6C8E3A")]
        vals = [[(v if v is not None else 0) for v in se["values"]] for se in s["series"]]
        if not s["labels"] or not any(any(v) for v in vals):
            d.add(String(10, h / 2, "No data", fontName="Helvetica", fontSize=10, fillColor=C(MUTED)))
            return d
        if s["kind"] == "line":
            c = HorizontalLineChart()
            c.x, c.y, c.width, c.height = 40, 30, W - 200, h - 50
            c.data = vals
            c.categoryAxis.categoryNames = s["labels"]
            for i in range(len(vals)):
                c.lines[i].strokeColor = palette[i % 4]
                c.lines[i].strokeWidth = 2.2
                from reportlab.graphics.widgets.markers import makeMarker
                c.lines[i].symbol = makeMarker("FilledCircle", size=4, fillColor=palette[i % 4])
        elif s["kind"] == "bar":
            c = VerticalBarChart()
            c.x, c.y, c.width, c.height = 40, 30, W - 200, h - 50
            c.data = vals
            c.categoryAxis.categoryNames = s["labels"]
            c.barSpacing = 2
            for i in range(len(vals)):
                c.bars[i].fillColor = palette[i % 4]
                c.bars[i].strokeColor = None
            if s.get("colors") and len(vals) == 1:
                for j, col in enumerate(s["colors"]):
                    c.bars[(0, j)].fillColor = C(col)
        else:
            c = HorizontalBarChart()
            c.x, c.y, c.width, c.height = 150, 15, W - 300, h - 25
            c.data = [list(reversed(v)) for v in vals]
            c.categoryAxis.categoryNames = list(reversed(s["labels"]))
            c.bars[0].fillColor = C(BROWN2)
            c.bars[0].strokeColor = None
            c.valueAxis.valueMin = 0
        c.categoryAxis.labels.fontName = "Helvetica"
        c.categoryAxis.labels.fontSize = 7.5
        c.valueAxis.labels.fontSize = 7.5
        c.valueAxis.strokeColor = C(LINE)
        c.categoryAxis.strokeColor = C(LINE)
        c.valueAxis.gridStrokeColor = C(LINE)
        c.valueAxis.visibleGrid = True
        d.add(c)
        if len(vals) > 1:
            lg = Legend()
            lg.x, lg.y = W - 145, h - 20
            lg.fontName, lg.fontSize = "Helvetica", 8
            lg.colorNamePairs = [(palette[i % 4], se["name"]) for i, se in enumerate(s["series"])]
            d.add(lg)
        return d

    for s in rep["sections"]:
        t = s["type"]
        if t == "pagebreak":
            story.append(PageBreak())
        elif t == "heading":
            story += [Paragraph(esc(s["title"]), ParagraphStyle("hh", parent=st_title, fontSize=16)), Paragraph(esc(s.get("sub", "")), st_sub),
                      Spacer(1, 3 * mm)]
        elif t == "kpis":
            block = ([Paragraph(esc(s["title"]), st_h)] if s.get("title") else []) + [kpis(s["items"]), Spacer(1, 4 * mm)]
            story.append(KeepTogether(block))
        elif t == "table":
            story += [Paragraph(esc(s["title"]), st_h), table(s) if s["rows"] else Paragraph("Nothing to show.", st_sub), Spacer(1, 4 * mm)]
        elif t == "chart":
            story.append(KeepTogether([Paragraph(esc(s["title"]), st_h), chart(s), Spacer(1, 3 * mm)]))
        elif t == "text":
            story += [Paragraph(esc(s["title"]), st_h)] + [Paragraph("• " + esc(line), st_txt) for line in s["lines"]] + [Spacer(1, 3 * mm)]
    if rep.get("sources"):
        story += [Paragraph("Sources", st_h)] + [Paragraph(esc(x), st_sub) for x in rep["sources"]]

    def deco(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(C(BROWN))
        canvas.rect(0, page[1] - 9 * mm, page[0], 9 * mm, stroke=0, fill=1)
        canvas.setFillColor(C(GOLD))
        canvas.rect(0, page[1] - 9.8 * mm, page[0], 0.8 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawString(15 * mm, page[1] - 6 * mm, "PNO  ·  People & Performance")
        canvas.setFont("Helvetica", 8.5)
        canvas.drawRightString(page[0] - 15 * mm, page[1] - 6 * mm, f"{rep['scope']}  ·  {rep['month_label']}")
        canvas.setFillColor(C(MUTED))
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(15 * mm, 8 * mm, f"Confidential – HR  ·  Generated {rep['generated']}" + ("  ·  names hidden" if rep["hide"] else ""))
        canvas.drawRightString(page[0] - 15 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=page, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=16 * mm, bottomMargin=14 * mm,
                            title=f"PNO {rep['title']}", author="PNO")
    doc.build(story, onFirstPage=deco, onLaterPages=deco)


# ------------------------------------------------------------------------------------------------ PowerPoint
def render_pptx(rep: dict, path: str) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    def rgb(h):
        return RGBColor.from_string(h.lstrip("#"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    SW, SH = prs.slide_width, prs.slide_height

    def box(sl, x, y, w, h, color, line=None):
        shp = sl.shapes.add_shape(1, x, y, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(color)
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def text(sl, x, y, w, h, s, size=14, bold=False, color=INK, align=None):
        tb = sl.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s
        p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(size), bold, rgb(color), "Segoe UI"
        if align:
            p.alignment = align
        return tb

    def frame(title, sub=None):
        sl = prs.slides.add_slide(blank)
        box(sl, 0, 0, SW, Inches(0.9), BROWN)
        box(sl, 0, Inches(0.9), SW, Inches(0.05), GOLD)
        text(sl, Inches(0.5), Inches(0.18), Inches(9), Inches(0.6), title, 24, True, "#FFFFFF")
        text(sl, Inches(8.5), Inches(0.28), Inches(4.4), Inches(0.5), f"{rep['scope']} · {rep['month_label']}", 11, False, "#F2E8DC", PP_ALIGN.RIGHT)
        if sub:
            text(sl, Inches(0.5), Inches(1.05), Inches(12), Inches(0.4), sub, 12, False, MUTED)
        text(sl, Inches(0.5), SH - Inches(0.45), Inches(8), Inches(0.3), "PNO · Confidential – HR" + (" · names hidden" if rep["hide"] else ""), 9, False, MUTED)
        return sl

    # title slide
    sl = prs.slides.add_slide(blank)
    box(sl, 0, 0, SW, SH, BROWN)
    box(sl, Inches(0.8), Inches(3.55), Inches(1.4), Inches(0.08), GOLD)
    text(sl, Inches(0.8), Inches(1.6), Inches(11), Inches(1.0), "PNO · People & Performance", 18, False, "#D7A23A")
    text(sl, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.3), rep["title"], 40, True, "#FFFFFF")
    text(sl, Inches(0.8), Inches(3.8), Inches(11.5), Inches(0.8), rep.get("subtitle") or f"{rep['scope']} · {rep['month_label']}", 18, False, "#F2E8DC")
    text(sl, Inches(0.8), Inches(6.6), Inches(11.5), Inches(0.5), f"Generated {rep['generated']} · Confidential – HR", 11, False, "#BFAE9C")

    pending_heading = None
    for s in rep["sections"]:
        t = s["type"]
        if t == "heading":
            pending_heading = (s["title"], s.get("sub"))
            continue
        if t == "pagebreak":
            continue
        if t == "kpis":
            title = s.get("title") or (pending_heading[0] if pending_heading else "Key numbers")
            sl = frame(title, pending_heading[1] if pending_heading else None)
            pending_heading = None
            n = len(s["items"])
            gap = Inches(0.25)
            w = (SW - Inches(1.0) - gap * (n - 1)) / n
            for i, it in enumerate(s["items"]):
                x = Inches(0.5) + i * (w + gap)
                y = Inches(1.9)
                box(sl, x, y, w, Inches(2.6), "#FFFFFF", LINE)
                box(sl, x, y, w, Inches(0.07), GOLD)
                text(sl, x + Inches(0.2), y + Inches(0.25), w - Inches(0.4), Inches(0.4), it["label"].upper(), 11, True, MUTED)
                text(sl, x + Inches(0.2), y + Inches(0.75), w - Inches(0.4), Inches(1.0), str(it["value"]), 34, True, BROWN)
                text(sl, x + Inches(0.2), y + Inches(1.75), w - Inches(0.4), Inches(0.7), it.get("sub", ""), 11, False, MUTED)
        elif t == "chart":
            sl = frame(s["title"])
            cd = CategoryChartData()
            cd.categories = s["labels"]
            for se in s["series"]:
                cd.add_series(se["name"], [v if v is not None else None for v in se["values"]])
            ct = {"line": XL_CHART_TYPE.LINE_MARKERS, "bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "hbar": XL_CHART_TYPE.BAR_CLUSTERED}[s["kind"]]
            gf = sl.shapes.add_chart(ct, Inches(0.6), Inches(1.4), SW - Inches(1.2), SH - Inches(2.2), cd)
            ch = gf.chart
            ch.has_legend = len(s["series"]) > 1
            if ch.has_legend:
                ch.legend.position = XL_LEGEND_POSITION.BOTTOM
                ch.legend.include_in_layout = False
            palette = [BROWN2, "#3B6EA5", GOLD, "#6C8E3A"]
            for i, ser in enumerate(ch.series):
                if s["kind"] == "line":
                    ser.format.line.color.rgb = rgb(palette[i % 4])
                    ser.format.line.width = Pt(2.5)
                    ser.smooth = False
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = rgb(palette[i % 4])
                    if s.get("colors") and len(s["series"]) == 1:
                        for j, col in enumerate(s["colors"]):
                            pt = ser.points[j]
                            pt.format.fill.solid()
                            pt.format.fill.fore_color.rgb = rgb(col)
            if s["kind"] == "hbar":
                ch.plots[0].categories  # noqa: B018
                ch.category_axis.reverse_order = True
            ch.font.size = Pt(11)
            ch.font.name = "Segoe UI"
        elif t == "table":
            cols = s["columns"][:9]
            rows = s["rows"]
            per = 12
            chunks = [rows[i:i + per] for i in range(0, max(len(rows), 1), per)][:8]
            for k, chunk in enumerate(chunks):
                sl = frame(s["title"] + (f" ({k + 1}/{len(chunks)})" if len(chunks) > 1 else ""),
                           pending_heading[1] if pending_heading and k == 0 else None)
                tbl = sl.shapes.add_table(len(chunk) + 1, len(cols), Inches(0.5), Inches(1.45), SW - Inches(1.0),
                                          Inches(0.38) * (len(chunk) + 1)).table
                for j, (lbl, key, fm) in enumerate(cols):
                    c = tbl.cell(0, j)
                    c.text = lbl
                    c.fill.solid()
                    c.fill.fore_color.rgb = rgb(BROWN)
                    para = c.text_frame.paragraphs[0]
                    para.font.size, para.font.bold, para.font.color.rgb = Pt(11), True, rgb("#FFFFFF")
                for i, row in enumerate(chunk, 1):
                    for j, (lbl, key, fm) in enumerate(cols):
                        c = tbl.cell(i, j)
                        v = row.get(key)
                        c.text = _f(v, fm) if fm else (str(v) if v not in (None, "") else "—")
                        c.fill.solid()
                        c.fill.fore_color.rgb = rgb("#FFFFFF" if i % 2 else BG)
                        para = c.text_frame.paragraphs[0]
                        para.font.size, para.font.color.rgb = Pt(10.5), rgb(INK)
            if len(rows) > per * 8:
                text(sl, Inches(0.5), SH - Inches(0.8), Inches(12), Inches(0.3),
                     f"Showing {per * 8} of {len(rows)} rows · export to Excel for the full list.", 10, False, MUTED)
            pending_heading = None
        elif t == "text":
            sl = frame(s["title"])
            tb = sl.shapes.add_textbox(Inches(0.7), Inches(1.5), SW - Inches(1.4), SH - Inches(2.3))
            tf = tb.text_frame
            tf.word_wrap = True
            for i, line in enumerate(s["lines"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = "•  " + line
                p.font.size, p.font.color.rgb, p.font.name = Pt(18), rgb(INK), "Segoe UI"
                p.space_after = Pt(12)
    if rep.get("sources"):
        sl = frame("Sources")
        tb = sl.shapes.add_textbox(Inches(0.7), Inches(1.5), SW - Inches(1.4), SH - Inches(2.3))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(rep["sources"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size, p.font.color.rgb = Pt(13), rgb(MUTED)
    _ = Emu
    prs.save(path)
