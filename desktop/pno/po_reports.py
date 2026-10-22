"""Board pack (landscape PDF), leader one-pagers and the Excel data workbook."""

from __future__ import annotations

import re
from datetime import datetime
from html import escape
from typing import Any

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from . import privacy
from .insights import WEIGHT_LABELS, Insights, Leader, tenure_text
from .po_format import BAD, GOOD, NEUTRAL, WARN, leader_name, money, number, pct, pp, score_text, tone

BROWN = colors.HexColor("#31261D")
GOLD = colors.HexColor("#B5975A")
CREAM = colors.HexColor("#F6F1E7")
LINE = colors.HexColor("#E7E2D9")

H1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=22, textColor=BROWN, leading=26, spaceAfter=4)
H2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=14, textColor=BROWN, leading=18, spaceBefore=6, spaceAfter=6)
BODY = ParagraphStyle("body", fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#28231F"))
SMALL = ParagraphStyle("small", parent=BODY, fontSize=7.5, leading=9.5, textColor=colors.HexColor("#756E68"))
CELL = ParagraphStyle("cell", parent=BODY, fontSize=8, leading=10)
CELL_B = ParagraphStyle("cellb", parent=CELL, fontName="Helvetica-Bold")
KPI_V = ParagraphStyle("kpiv", fontName="Helvetica-Bold", fontSize=18, leading=21, textColor=BROWN)
KPI_L = ParagraphStyle("kpil", fontName="Helvetica-Bold", fontSize=7, leading=9, textColor=colors.HexColor("#756E68"))


def _p(text: Any, style: ParagraphStyle = CELL, color: str | None = None) -> Paragraph:
    body = escape(str(text))
    if color:
        body = f"<font color='{color}'>{body}</font>"
    return Paragraph(body, style)


def _table(header: list[str], rows: list[list[Any]], widths: list[float] | None = None, fills: dict[tuple[int, int], str] | None = None) -> Table:
    data = [[_p(h, CELL_B) for h in header]] + rows
    table = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), CREAM),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, GOLD),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, LINE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for (c, r), color in (fills or {}).items():
        style.append(("BACKGROUND", (c, r + 1), (c, r + 1), colors.HexColor(color)))
    table.setStyle(TableStyle(style))
    return table


def _footer(canvas: Any, doc: Any) -> None:
    canvas.saveState()
    canvas.setFillColor(BROWN)
    canvas.rect(0, doc.pagesize[1] - 9 * mm, doc.pagesize[0], 9 * mm, stroke=0, fill=1)
    canvas.setFillColor(GOLD)
    canvas.setFont("Helvetica-Bold", 9)
    canvas.drawString(12 * mm, doc.pagesize[1] - 6 * mm, "PNO · CARREFOUR PAKISTAN · PEOPLE & ORGANISATION")
    canvas.setFillColor(colors.white)
    canvas.drawRightString(doc.pagesize[0] - 12 * mm, doc.pagesize[1] - 6 * mm, getattr(doc, "pno_title", ""))
    canvas.setFillColor(colors.HexColor("#756E68"))
    canvas.setFont("Helvetica", 7)
    note = "Confidential · generated offline from imported P&O workbook tabs"
    if privacy.demo_mode():
        note += " · presentation mode: names hidden"
    canvas.drawString(12 * mm, 7 * mm, note)
    canvas.drawRightString(doc.pagesize[0] - 12 * mm, 7 * mm, f"Page {doc.page}")
    canvas.restoreState()


def _doc(path: str, title: str) -> SimpleDocTemplate:
    doc = SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=15 * mm, bottomMargin=13 * mm, title=title, author="PNO")
    doc.pno_title = title
    return doc


def _kpis(items: list[tuple[str, str, str]]) -> Table:
    cells = [[Paragraph(escape(label.upper()), KPI_L), Paragraph(escape(value), KPI_V), Paragraph(escape(note), SMALL)] for label, value, note in items]
    table = Table([cells], colWidths=[(landscape(A4)[0] - 24 * mm) / len(items)] * len(items))
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), CREAM), ("BOX", (0, 0), (-1, -1), 0.5, LINE),
        ("INNERGRID", (0, 0), (-1, -1), 2, colors.white), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    return table


def _heat(value: float | None, good: float, bad: float, higher: bool = True) -> str:
    if value is None:
        return "#F4F2EE"
    v, g, b = (value, good, bad) if higher else (-value, -good, -bad)
    return "#CFE6D5" if v >= g else ("#F3CFCB" if v <= b else "#F7E7C2")


def _leader_flow(insights: Insights, leader: Leader) -> list[Any]:
    flow: list[Any] = []
    flow.append(Paragraph(escape(leader_name(leader)), H1))
    flow.append(Paragraph(escape(f"{leader.designation or 'Leader'} · {leader.store or '—'} · leads {leader.section_text}"), BODY))
    flow.append(Spacer(1, 6))
    box = insights.talent_box(leader)
    flow.append(_kpis([
        ("P&O index", score_text(leader.score), "0–100, weighted"),
        ("Rating", "—" if leader.latest_rating is None else f"{leader.latest_rating:g}", leader.rating_label),
        ("In role", tenure_text(leader.months_in_role), f"service {tenure_text(leader.months_service)}"),
        ("Talent grid", insights.TALENT_NAMES.get(box, "—") if box else "—", f"grade {leader.grade or '—'}"),
    ]))
    flow.append(Spacer(1, 8))
    rows, fills = [], {}
    k = leader.kpis
    for section, s in k.get("sales", []):
        fills[(2, len(rows))] = _heat(s["vs_budget_pct"], 100, 92)
        rows.append([_p(f"Sales · {section}"), _p(money(s["actual"])), _p(f"{pct(s['vs_budget_pct'], digits=0)} of budget"), _p(f"{pct(s['growth_pct'], True)} vs LY")])
    for section, p in k.get("productivity", []):
        fills[(2, len(rows))] = _heat(p["vs_target_pct"], 0, -10)
        level = f" (at {p['level']} level)" if p.get("fallback") else ""
        rows.append([_p(f"Productivity · {section}{level}"), _p(number(p["actual"])), _p(f"{pct(p['vs_target_pct'], True, 0)} vs target {number(p['target'])}"), _p(f"{pct(p['vs_ly_pct'], True, 0)} vs LY")])
    for section, sh in k.get("shrink", []):
        fills[(2, len(rows))] = _heat(sh["shrink_variance_pct"], 0, -1)
        rows.append([_p(f"Shrinkage · {section}"), _p(pct(sh["shrink_actual_pct"], digits=2)), _p(f"budget {pct(sh['shrink_budget_pct'], digits=2)} ({pp(sh['shrink_variance_pct'])})"), _p(f"stock days {number(sh['stock_days_actual'])} vs {number(sh['stock_days_budget'])}")])
    if rows:
        flow.append(Paragraph("Business results of their area", H2))
        flow.append(_table(["Measure", "Value", "Against plan", "Context"], rows, [85 * mm, 40 * mm, 75 * mm, 70 * mm], fills))
    people = []
    if k.get("leavers") is not None:
        people.append(f"{k['leavers']} leaver(s) in their section(s) in the last 12 months")
    if k.get("da") is not None:
        people.append(f"{k['da']} disciplinary case(s) recorded for their section(s)")
    if leader.shifts:
        people.append("Roster this month: " + ", ".join(f"{leader.shifts[s]} {s.lower()}" for s in ("Morning", "Mid", "Evening", "Off", "Leave") if leader.shifts[s]))
    for f in insights.bench():
        if f["leader"] is leader:
            people.append(f"Flag ({f['severity']}): {f['reason']}")
    if people:
        flow.append(Paragraph("People", H2))
        for line in people:
            flow.append(Paragraph("• " + escape(line), BODY))
    if leader.components:
        flow.append(Paragraph("How the index is built", H2))
        flow.append(_table(["Measure", "Score /100", "Weight"], [
            [_p(WEIGHT_LABELS[c]), _p(f"{v:.0f}"), _p(f"{insights.weights.get(c, 0):g}")] for c, v in leader.components.items()
        ], [90 * mm, 40 * mm, 40 * mm]))
        flow.append(Paragraph("Each measure is scaled 0–100 and weighted; missing measures are excluded, never treated as zero. A decision aid, not an automated decision.", SMALL))
    return flow


def export_leader_pdf(insights: Insights, leader: Leader, path: str) -> str:
    doc = _doc(path, f"Leader 360° · {leader_name(leader)}")
    doc.build(_leader_flow(insights, leader), onFirstPage=_footer, onLaterPages=_footer)
    return path


def export_board_pack(insights: Insights, path: str, one_pagers: int = 20) -> str:
    h = insights.headline()
    doc = _doc(path, f"Board pack · {h['as_of']:%B %Y}")
    width = landscape(A4)[0] - 24 * mm
    flow: list[Any] = []
    # cover + headline
    flow.append(Spacer(1, 20 * mm))
    flow.append(Paragraph("Country P&amp;O Review", ParagraphStyle("cover", parent=H1, fontSize=34, leading=40)))
    flow.append(Paragraph(escape(f"Carrefour Pakistan · data as of {h['as_of']:%d %B %Y} · prepared {datetime.now():%d %b %Y}"), BODY))
    flow.append(Spacer(1, 12 * mm))
    flow.append(_kpis([
        ("Sales MTD", money(h["sales"]), f"{pct(h['sales_vs_budget_pct'])} of budget"),
        ("Growth vs LY", pct(h["growth_pct"], True), "country"),
        ("Productivity vs target", pct(h["productivity_vs_target_pct"], True), "avg across stores"),
        ("Stores on track", f"{h['stores_on_track']} / {h['stores']}", f"{h['stores_at_risk']} at risk"),
    ]))
    flow.append(Spacer(1, 4))
    flow.append(_kpis([
        ("Section leaders", str(h["leaders"]), f"avg rating {number(h['avg_rating'], 2)}"),
        ("Leavers · 12 months", number(h["leavers_12m"]), f"{pct(h['voluntary_pct'], digits=0)} voluntary"),
        ("Disciplinary cases", number(h["da_cases"]), "DA tab"),
        ("Open alerts", str(len(insights.alerts())), f"{sum(1 for a in insights.alerts() if a['severity'] == 'High')} high priority"),
    ]))
    flow.append(Spacer(1, 8))
    top = insights.alerts(6)
    if top:
        flow.append(Paragraph("Top priorities", H2))
        for a in top:
            flow.append(Paragraph(f"<font color='{BAD if a['severity'] == 'High' else WARN}'><b>{a['severity']}</b></font> · {escape(a['area'])} · <b>{escape(a['store'])}</b> — {escape(a['message'])}", BODY))
    flow.append(PageBreak())
    # league table
    flow.append(Paragraph("Store league table", H2))
    rows, fills = [], {}
    for i, c in enumerate(insights.store_scorecard()):
        fills[(1, i)] = {"On track": "#CFE6D5", "Watch": "#F7E7C2", "At risk": "#F3CFCB"}.get(c["status"], "#F4F2EE")
        fills[(3, i)] = _heat(c["sales_vs_budget_pct"], 100, 92)
        fills[(5, i)] = _heat(c["productivity_vs_target_pct"], 0, -10)
        fills[(6, i)] = _heat(c["shrink_variance_avg"], 0, -1)
        rows.append([_p(c["store"], CELL_B), _p(c["status"]), _p(money(c["sales"])), _p(pct(c["sales_vs_budget_pct"], digits=0)),
                     _p(pct(c["growth_pct"], True)), _p(pct(c["productivity_vs_target_pct"], True, 0)), _p(pp(c["shrink_variance_avg"])),
                     _p(pct(c["oos_pct"])), _p(pct(c["margin_pct"])), _p(number(c["leavers_12m"])), _p(number(c["da_cases"])),
                     _p(score_text(c["avg_leader_score"]))])
    flow.append(_table(["Store", "Status", "Sales", "vs Bud.", "Growth", "Prod. vs tgt", "Shrink", "OOS", "Margin", "Leavers", "DA", "Leader idx"],
                       rows, [44 * mm, 18 * mm, 22 * mm, 17 * mm, 18 * mm, 21 * mm, 19 * mm, 15 * mm, 17 * mm, 17 * mm, 12 * mm, 18 * mm], fills))
    flow.append(Paragraph("Status combines sales vs budget, productivity vs target and shrinkage vs budget.", SMALL))
    flow.append(PageBreak())
    # heatmaps
    for title, metric, good, bad, higher, fmt in (
        ("Sales vs budget by section (%)", "sales_vs_budget", 100, 92, True, "{:.0f}"),
        ("Shrinkage vs budget by section (pp)", "shrink_variance", 0, -1, True, "{:+.1f}"),
        ("Productivity vs target by section (%)", "productivity_vs_target", 0, -10, True, "{:+.0f}"),
    ):
        stores, sections, values = insights.section_matrix(metric)
        if not sections:
            continue
        rows, fills = [], {}
        for r, store_name in enumerate(stores):
            row = [_p(store_name, CELL_B)]
            for c, section in enumerate(sections):
                v = values.get((store_name, section))
                fills[(c + 1, r)] = _heat(v, good, bad, higher)
                row.append(_p("" if v is None else fmt.format(v)))
            rows.append(row)
        col = (width - 44 * mm) / len(sections)
        flow.append(KeepTogether([Paragraph(title, H2), _table(["Store"] + [s.replace("Fruits & Vegetables", "F&V").replace("Light Household", "LHH").replace("Heavy Household", "HHH") for s in sections], rows, [44 * mm] + [col] * len(sections), fills)]))
        flow.append(Spacer(1, 6))
    flow.append(PageBreak())
    # people
    flow.append(Paragraph("People: leaders, talent and bench strength", H2))
    leaders = [l for l in insights.leaders() if l.score is not None]
    def leader_rows(items: list[Leader]) -> list[list[Any]]:
        return [[_p(leader_name(l), CELL_B), _p(l.store), _p(l.section_text), _p(l.designation), _p(score_text(l.score), color=tone(l.score, 70, 45)),
                 _p("—" if l.latest_rating is None else f"{l.latest_rating:g}"), _p(tenure_text(l.months_in_role))] for l in items]
    head = ["Leader", "Store", "Leads", "Role", "Index", "Rating", "In role"]
    widths = [55 * mm, 45 * mm, 50 * mm, 35 * mm, 18 * mm, 18 * mm, 25 * mm]
    if leaders:
        flow.append(Paragraph("Top 10 by P&amp;O index", BODY))
        flow.append(_table(head, leader_rows(leaders[:10]), widths))
        flow.append(Spacer(1, 6))
        flow.append(Paragraph("Bottom 10 by P&amp;O index", BODY))
        flow.append(_table(head, leader_rows(leaders[-10:][::-1]), widths))
    bench = insights.bench()
    if bench:
        flow.append(PageBreak())
        flow.append(Paragraph("Bench strength: roles at risk", H2))
        flow.append(_table(["Priority", "Leader", "Store", "Leads", "Why flagged"], [
            [_p(f["severity"], color=BAD if f["severity"] == "High" else WARN), _p(leader_name(f["leader"]), CELL_B), _p(f["leader"].store),
             _p(f["leader"].section_text), _p(f["reason"])] for f in bench], [20 * mm, 55 * mm, 45 * mm, 50 * mm, 95 * mm]))
    if insights.has("po_turnover"):
        t = insights.turnover_summary()
        flow.append(PageBreak())
        flow.append(Paragraph("Turnover · last 12 months", H2))
        flow.append(_kpis([("Leavers", str(t["total"]), f"{t['all_time']} on file"), ("Voluntary", pct(t["voluntary_pct"], digits=0), "resignations"),
                           ("Top reason", t["by_reason"][0][0] if t["by_reason"] else "—", ""), ("Top hotspot", f"{t['hotspots'][0][0]}" if t["hotspots"] else "—", f"{t['hotspots'][0][1]}: {t['hotspots'][0][2]}" if t["hotspots"] else "")]))
        flow.append(Spacer(1, 6))
        half = width / 2 - 4 * mm
        left = _table(["Reason", "Leavers"], [[_p(r), _p(str(n))] for r, n in t["by_reason"][:10]], [half - 25 * mm, 25 * mm])
        right = _table(["Store · section", "Leavers"], [[_p(f"{s} · {sec}"), _p(str(n))] for s, sec, n in t["hotspots"][:10]], [half - 25 * mm, 25 * mm])
        flow.append(Table([[left, right]], colWidths=[half + 4 * mm, half + 4 * mm], style=[("VALIGN", (0, 0), (-1, -1), "TOP")]))
    # one-pagers
    for leader in leaders[:one_pagers]:
        flow.append(PageBreak())
        flow.extend(_leader_flow(insights, leader))
    doc.build(flow, onFirstPage=_footer, onLaterPages=_footer)
    return path


def export_workbook(insights: Insights, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="31261D")
    header_font = Font(color="FFFFFF", bold=True)

    def sheet(title: str, headers: list[str], rows: list[list[Any]]) -> None:
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append([_safe(v) for v in row])
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(48, len(h) + 6))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    cards = insights.store_scorecard()
    sheet("Store scorecard", ["Store", "Status", "Sales", "Budget", "Sales vs budget %", "Growth %", "Margin %", "Waste %", "OOS %",
                              "Productivity", "Productivity target", "Productivity vs target %", "Avg shrink variance pp",
                              "Sections over stock days", "NPS", "Leavers 12m", "DA cases", "Leaders", "Avg rating", "Avg leader index"],
          [[c["store"], c["status"], c["sales"], c["budget"], c["sales_vs_budget_pct"], c["growth_pct"], c["margin_pct"], c["waste_pct"], c["oos_pct"],
            c["productivity"], c["productivity_target"], c["productivity_vs_target_pct"], c["shrink_variance_avg"], c["sections_over_stock_days"],
            c["nps"], c["leavers_12m"], c["da_cases"], c["leaders"], c["avg_rating"], c["avg_leader_score"]] for c in cards])
    sheet("Alerts", ["Severity", "Area", "Store", "Message"], [[a["severity"], a["area"], a["store"], a["message"]] for a in insights.alerts()])
    lead_rows = []
    for l in insights.leaders():
        lead_rows.append([leader_name(l), privacy.employee_id(l.employee_id), l.store, l.section_text, l.designation, l.grade, l.score, l.business,
                          l.latest_rating, l.rating_label, l.months_in_role, l.months_service, privacy.name(l.manager_name) if l.manager_name else "",
                          *[l.components.get(k) for k in WEIGHT_LABELS]])
    sheet("Leaders", ["Leader", "Employee number", "Store", "Leads", "Designation", "Grade", "P&O index", "Business results", "Latest rating",
                      "Rating label", "Months in role", "Months of service", "Reports to", *[f"{v} (0-100)" for v in WEIGHT_LABELS.values()]], lead_rows)
    sheet("Bench strength", ["Priority", "Leader", "Store", "Leads", "Reason"],
          [[f["severity"], leader_name(f["leader"]), f["leader"].store, f["leader"].section_text, f["reason"]] for f in insights.bench()])
    for title, metric in (("Heat · sales vs budget", "sales_vs_budget"), ("Heat · shrink vs budget", "shrink_variance"), ("Heat · productivity", "productivity_vs_target")):
        stores, sections, values = insights.section_matrix(metric)
        sheet(title, ["Store", *sections], [[s, *[values.get((s, sec)) for sec in sections]] for s in stores])
    if insights.has("po_turnover"):
        t = insights.turnover_summary()
        sheet("Turnover hotspots", ["Store", "Section", "Leavers 12m"], [list(x) for x in t["hotspots"]])
        sheet("Turnover reasons", ["Reason", "Leavers 12m"], [list(x) for x in t["by_reason"]])
    r = insights.roster_summary()
    if r.get("gaps"):
        sheet("Roster gaps", ["Store", "Department", "Date", "Leaders off"], [[g["store"], g["department"], g["date"], g["leaders"]] for g in r["gaps"]])
    wb.save(path)
    return path


def _safe(value: Any) -> Any:
    if isinstance(value, float):
        return round(value, 2)
    if isinstance(value, str) and re.match(r"^[=+\-@]", value):
        return "'" + value
    return value
