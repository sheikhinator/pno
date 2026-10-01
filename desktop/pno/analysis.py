"""Explainable, local-only analysis and report exports for PNO data."""

from __future__ import annotations

import calendar
import csv
import json
import math
import re
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Iterable


METRIC_NAMES = (
    "actual",
    "budget",
    "last_year",
    "budget_achievement_pct",
    "sales_growth_pct",
    "oos_value",
    "oos_pct",
    "attendance_rate_pct",
    "hours",
    "employee_count",
    "direct_reports",
    "productivity",
)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        raw = str(value).strip().replace(",", "").replace("%", "")
        if not raw:
            return None
        try:
            result = float(raw)
        except (TypeError, ValueError):
            return None
    return result if math.isfinite(result) else None


def _first(row: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        value = row.get(key)
        if value is not None and _text(value):
            return value
    return None


def _scope_value(row: dict[str, Any], dimension: str) -> str:
    aliases = {
        "country": ("country", "country_name"),
        "store": ("store", "store_name"),
        "department": ("department", "dept"),
        "section": ("section",),
        "employee_id": ("employee_id", "employee_code", "staff_id"),
    }
    value = _first(row, *aliases[dimension])
    if value is None and dimension == "country" and isinstance(row.get("raw"), dict):
        for key, raw_value in row["raw"].items():
            normalized = re.sub(r"[\W_]+", "", str(key).casefold())
            if normalized in {"country", "countryname"} and _text(raw_value):
                return _text(raw_value)
    return _text(value)


def _period_date(value: Any) -> date | None:
    raw = _text(value)
    if not raw:
        return None
    for candidate in (raw[:10], raw[:7] + "-01"):
        try:
            return date.fromisoformat(candidate)
        except ValueError:
            pass
    return None


def _row_period(row: dict[str, Any]) -> str:
    return _text(row.get("date") or row.get("_period"))


def _snapshot_period(row: dict[str, Any]) -> str:
    return _snapshot_stamp(row)[:10]


def _snapshot_stamp(row: dict[str, Any]) -> str:
    batch_period = _text(row.get("_period"))
    record_date = _text(row.get("date"))
    if len(batch_period) >= 10 and _period_date(batch_period):
        return batch_period
    return record_date or batch_period


def _observation_period(row: dict[str, Any], kind: str) -> str:
    if kind == "sales_mtd":
        return _snapshot_period(row)
    return _text(row.get("date") or row.get("_period"))[:10]


def _scope_matches(row: dict[str, Any], filters: dict[str, Any]) -> bool:
    for key in ("country", "store", "department", "section"):
        requested = _text(filters.get(key))
        if requested and _scope_value(row, key).casefold() != requested.casefold():
            return False
    return True


def _remove_overlapping_summaries(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Keep summary rows in preference to their descendant detail rows."""
    dimensions = ("country", "store", "department", "section", "employee_id")

    def values(row: dict[str, Any]) -> dict[str, str]:
        return {key: _scope_value(row, key).casefold() for key in dimensions}

    indexed = [(row, values(row)) for row in rows]
    result: list[dict[str, Any]] = []
    for row, detail in indexed:
        is_covered = False
        for other, summary in indexed:
            if other is row:
                continue
            summary_dimensions = {key for key, value in summary.items() if value}
            detail_dimensions = {key for key, value in detail.items() if value}
            if not summary_dimensions or len(summary_dimensions) >= len(detail_dimensions):
                continue
            if all(detail[key] == value for key, value in summary.items() if value):
                # A country/store/dept/section total covers its descendants.
                # Require a geographic scope to avoid generic blank rows hiding
                # otherwise disjoint records.
                if (summary.get("country") or summary.get("store")):
                    is_covered = True
                    break
        if not is_covered:
            result.append(row)
    return result


def _collapse_by_observation(rows: list[dict[str, Any]], kind: str) -> list[dict[str, Any]]:
    """Collapse summary/detail overlap only within the same observation date."""
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        period = _observation_period(row, kind)
        grouped[period].append(row)
    collapsed: list[dict[str, Any]] = []
    for period in sorted(grouped):
        collapsed.extend(_remove_overlapping_summaries(grouped[period]))
    return collapsed


def _sum_complete(rows: list[dict[str, Any]], field: str) -> float | None:
    if not rows:
        return None
    values = [_number(row.get(field)) for row in rows]
    if any(value is None for value in values):
        return None
    return sum(value for value in values if value is not None)


def _percentage(numerator: float | None, denominator: float | None) -> float | None:
    if numerator is None or denominator is None or denominator == 0:
        return None
    return numerator * 100.0 / denominator


def _manager_role(employee: dict[str, Any]) -> str | None:
    title = _text(employee.get("designation")).casefold()
    if "general manager" in title or "store manager" in title or re.search(r"\bgm\b", title):
        return "GM"
    if any(phrase in title for phrase in ("department head", "department manager", "dept head", "dept manager")) or re.search(r"\bdh\b", title):
        return "DH"
    if any(phrase in title for phrase in ("section manager", "section mgr")) or re.search(r"\bsm\b", title):
        return "SM"
    return None


def _unique_employees(store: Any, as_of: str | None) -> tuple[list[dict[str, Any]], bool]:
    if not hasattr(store, "employees"):
        return [], False
    try:
        roster = store.employees(period=as_of)
    except TypeError:
        roster = store.employees(as_of)
    if not isinstance(roster, list):
        return [], False
    if not roster:
        # An empty result from the storage convenience method generally means
        # there is no active roster import, not a verified zero headcount.
        try:
            roster_records = store.records(kind="roster", period_to=as_of)
        except (AttributeError, TypeError):
            roster_records = []
        if not roster_records:
            return [], False
    unique: dict[tuple[str, str], dict[str, Any]] = {}
    newest_by_identity: dict[str, str] = {}
    for item in roster:
        if not isinstance(item, dict):
            continue
        employee_id = _text(item.get("employee_id") or item.get("employee_code"))
        identity = employee_id.casefold() if employee_id else _text(item.get("employee_name")).casefold()
        if identity:
            snapshot = _text(item.get("_period"))
            previous = newest_by_identity.get(identity)
            if previous is not None and snapshot < previous:
                continue
            if previous is None or snapshot > previous:
                unique = {
                    key: value
                    for key, value in unique.items()
                    if key[0] != identity
                }
                newest_by_identity[identity] = snapshot
            unique[(identity, _scope_value(item, "store").casefold())] = item
    return list(unique.values()), True


def _find_employee(roster: list[dict[str, Any]], query: Any) -> tuple[dict[str, Any] | None, list[dict[str, Any]]]:
    needle = _text(query).casefold()
    if not needle:
        return None, []
    id_matches = [
        employee for employee in roster
        if _text(employee.get("employee_id") or employee.get("employee_code")).casefold() == needle
    ]
    if id_matches:
        dated = [employee for employee in id_matches if _text(employee.get("_period"))]
        if dated:
            newest_period = max(_text(employee.get("_period")) for employee in dated)
            id_matches = [
                employee for employee in dated
                if _text(employee.get("_period")) == newest_period
            ]
    matches = id_matches or [
        employee for employee in roster
        if _text(employee.get("employee_name")).casefold() == needle
    ]
    # A repeated employee ID across stores is ambiguous unless filters narrowed it.
    unique: dict[tuple[str, str], dict[str, Any]] = {}
    for employee in matches:
        unique[(
            _text(employee.get("employee_id") or employee.get("employee_code")).casefold(),
            _scope_value(employee, "store").casefold(),
        )] = employee
    found = list(unique.values())
    return (found[0], found) if len(found) == 1 else (None, found)


def _employee_record_matches(row: dict[str, Any], employee: dict[str, Any]) -> bool:
    row_id = _text(row.get("employee_id") or row.get("employee_code")).casefold()
    employee_id = _text(employee.get("employee_id") or employee.get("employee_code")).casefold()
    if row_id and employee_id and row_id == employee_id:
        return True
    row_name = _text(row.get("employee_name")).casefold()
    employee_name = _text(employee.get("employee_name")).casefold()
    return bool(row_name and employee_name and row_name == employee_name)


def _roster_snapshot(store: Any, as_of: str | None) -> tuple[list[dict[str, Any]], bool]:
    roster, known = _unique_employees(store, as_of)
    if known:
        return roster, True
    # This fallback supports small Store-compatible implementations without an
    # employees() helper; normal desktop use resolves the storage API above.
    try:
        records = store.records(kind="roster", period_to=as_of)
    except (AttributeError, TypeError):
        return [], False
    if not isinstance(records, list):
        return [], False
    return records, bool(records)


def _source_records(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    sources: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in rows:
        if not isinstance(row, dict):
            continue
        kind = row.get("_kind")
        if not kind:
            continue
        source = {
            "import_id": row.get("_import_id"),
            "kind": kind,
            "period": row.get("_period"),
            "source": row.get("_source"),
            "sheet": row.get("_sheet"),
        }
        key = tuple(source.values())
        sources[key] = source
    return sorted(sources.values(), key=lambda item: (str(item.get("period") or ""), str(item.get("kind") or "")))


def _latest_mtd_snapshots(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Keep all rows from each scope's newest MTD snapshot in each month."""
    newest: dict[tuple[str, str, str, str, str, str], tuple[str, list[dict[str, Any]]]] = {}
    for row in rows:
        raw_period = _snapshot_stamp(row)
        parsed = _period_date(raw_period)
        month = parsed.strftime("%Y-%m") if parsed else raw_period[:7]
        scope = tuple(
            _scope_value(row, key).casefold()
            for key in ("country", "store", "department", "section", "employee_id")
        )
        key = (month, *scope)
        prior = newest.get(key)
        if prior is None or raw_period > prior[0]:
            newest[key] = (raw_period, [row])
        elif raw_period == prior[0]:
            prior[1].append(row)
    return [row for _, snapshot_rows in newest.values() for row in snapshot_rows]


def _collapse_mtd_hierarchy(
    rows: list[dict[str, Any]],
    warnings: list[str] | None = None,
) -> list[dict[str, Any]]:
    """Avoid summing overlapping MTD scopes across different snapshot dates."""
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        stamp = _snapshot_stamp(row)
        parsed = _period_date(stamp)
        month = parsed.strftime("%Y-%m") if parsed else stamp[:7]
        grouped[month].append(row)

    dimensions = ("country", "store", "department", "section", "employee_id")
    result: list[dict[str, Any]] = []
    for month, month_rows in sorted(grouped.items()):
        indexed = [
            (
                row,
                {key: _scope_value(row, key).casefold() for key in dimensions},
            )
            for row in month_rows
        ]
        covered: set[int] = set()
        has_older_summary = False
        for detail_index, (detail_row, detail) in enumerate(indexed):
            detail_dimensions = {key for key, value in detail.items() if value}
            for summary_index, (summary_row, summary) in enumerate(indexed):
                if summary_index == detail_index:
                    continue
                summary_dimensions = {key for key, value in summary.items() if value}
                if not summary_dimensions or len(summary_dimensions) >= len(detail_dimensions):
                    continue
                if not all(detail[key] == value for key, value in summary.items() if value):
                    continue
                # A summary covers descendants only when it has a geographic
                # anchor and every populated summary dimension matches.
                if not (summary.get("country") or summary.get("store")):
                    continue
                covered.add(detail_index)
                if _snapshot_stamp(summary_row) < _snapshot_stamp(detail_row):
                    has_older_summary = True
        if has_older_summary and warnings is not None:
            warnings.append(
                f"For MTD {month}, an older covering summary was retained as the coherent total; "
                "newer descendant snapshots were excluded because cross-snapshot coverage cannot be verified."
            )
        result.extend(row for index, (row, _) in enumerate(indexed) if index not in covered)
    return result


def _get_rows(store: Any, kind: str, filters: dict[str, Any]) -> list[dict[str, Any]]:
    # An employee filter is deliberately resolved from the as-of roster. Passing
    # it to records() would hide store/department/section totals used for managers.
    options = {
        "period_from": filters.get("period_from"),
        "period_to": filters.get("period_to"),
        "store": filters.get("store"),
        "department": filters.get("department"),
        "section": filters.get("section"),
    }
    rows = store.records(kind=kind, **options)
    return [
        row for row in rows
        if isinstance(row, dict) and _scope_matches(row, filters)
    ]


def _apply_manager_scope(
    employee: dict[str, Any],
    role: str,
    filters: dict[str, Any],
    warnings: list[str],
) -> dict[str, Any]:
    scoped = dict(filters)
    assigned = {
        "country": _scope_value(employee, "country"),
        "store": _scope_value(employee, "store"),
        "department": _scope_value(employee, "department"),
        "section": _scope_value(employee, "section"),
    }
    required = ("store",) if role == "GM" else ("store", "department") if role == "DH" else ("store", "department", "section")
    for key in required:
        assignment = assigned[key]
        requested = _text(filters.get(key))
        if not assignment:
            warnings.append(f"Manager assignment has no {key}; the assigned responsibility scope cannot be fully resolved.")
            scoped[key] = "\0missing-assignment"
            continue
        if requested and requested.casefold() != assignment.casefold():
            warnings.append(f"Requested {key} '{requested}' is outside the employee's as-of {role} responsibility.")
            scoped[key] = f"\0no-match:{requested}"
        else:
            scoped[key] = assignment
    # Narrower explicit filters remain useful overrides within responsibility.
    return scoped


def _as_of_employee(
    store: Any,
    employee_id: str,
    period: str,
    cache: dict[str, dict[str, Any] | None],
) -> dict[str, Any] | None:
    cache_key = period or "__latest__"
    if cache_key in cache:
        return cache[cache_key]

    candidates: list[dict[str, Any]] = []
    try:
        roster_rows = store.records(kind="roster", period_to=period or None)
    except (AttributeError, TypeError):
        roster_rows = []
    for row in roster_rows if isinstance(roster_rows, list) else []:
        if (
            isinstance(row, dict)
            and _text(row.get("employee_id") or row.get("employee_code")).casefold() == employee_id.casefold()
        ):
            candidates.append(row)

    if candidates:
        latest_period = max((_text(row.get("_period")) for row in candidates), default="")
        candidates = [row for row in candidates if _text(row.get("_period")) == latest_period]
    else:
        roster, known = _roster_snapshot(store, period or None)
        if known:
            candidates = [
                row for row in roster
                if _text(row.get("employee_id") or row.get("employee_code")).casefold() == employee_id.casefold()
            ]

    # Duplicate roster rows with the same as-of assignment are harmless; two
    # different assignments in one snapshot are ambiguous and fail closed.
    assignments = {
        (
            _manager_role(row),
            *(_scope_value(row, key).casefold() for key in ("country", "store", "department", "section")),
        )
        for row in candidates
    }
    resolved = candidates[0] if len(assignments) == 1 and candidates else None
    cache[cache_key] = resolved
    return resolved


def _manager_observation_matches(
    store: Any,
    row: dict[str, Any],
    employee_id: str,
    kind: str,
    filters: dict[str, Any],
    cache: dict[str, dict[str, Any] | None],
    warnings: list[str],
    *,
    allow_ordinary_employee: bool = False,
) -> bool:
    period = _observation_period(row, kind)
    manager = _as_of_employee(store, employee_id, period, cache)
    if manager is None:
        warnings.append(f"No unique as-of roster assignment for employee {employee_id} at {period or 'unknown period'}; the observation is excluded.")
        return False

    role = _manager_role(manager)
    if role is None:
        if allow_ordinary_employee:
            return _employee_record_matches(row, manager)
        warnings.append(f"Employee {employee_id} has no GM/DH/SM responsibility at {period or 'unknown period'}; sales are excluded.")
        return False

    assignment = {
        key: _scope_value(manager, key)
        for key in ("country", "store", "department", "section")
    }
    required = ("store",) if role == "GM" else ("store", "department") if role == "DH" else ("store", "department", "section")
    missing = [key for key in required if not assignment[key]]
    if missing:
        for key in missing:
            warnings.append(f"As-of {role} assignment for employee {employee_id} has no {key} at {period or 'unknown period'}; the observation fails closed.")
        return False

    if assignment["country"] and _scope_value(row, "country").casefold() != assignment["country"].casefold():
        return False
    for key in required:
        requested = _text(filters.get(key))
        if requested and requested.casefold() != assignment[key].casefold():
            warnings.append(f"Requested {key} '{requested}' is outside employee {employee_id}'s as-of {role} responsibility at {period}.")
            return False
        if _scope_value(row, key).casefold() != assignment[key].casefold():
            return False
    return True


def _filter_employee_observations(
    store: Any,
    rows: list[dict[str, Any]],
    employee: dict[str, Any],
    kind: str,
    filters: dict[str, Any],
    warnings: list[str],
    *,
    allow_ordinary_employee: bool = False,
) -> list[dict[str, Any]]:
    employee_id = _text(employee.get("employee_id") or employee.get("employee_code"))
    cache: dict[str, dict[str, Any] | None] = {}
    return [
        row for row in rows
        if _manager_observation_matches(
            store, row, employee_id, kind, filters, cache, warnings,
            allow_ordinary_employee=allow_ordinary_employee,
        )
    ]


def _attendance_metrics(rows: list[dict[str, Any]], warnings: list[str]) -> tuple[float | None, float | None]:
    exclude_terms = ("leave", "day off", "day-off", "off day", "not marked", "unmarked", "no mark")
    present_terms = ("present", "late")
    by_employee_day: dict[tuple[str, str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    unkeyed_rows: list[dict[str, Any]] = []
    for row in rows:
        observation_date = _text(row.get("date") or row.get("_period"))[:10]
        employee_id = _scope_value(row, "employee_id").casefold()
        if observation_date and employee_id:
            key = (
                observation_date,
                _scope_value(row, "country").casefold(),
                _scope_value(row, "store").casefold(),
                _scope_value(row, "department").casefold(),
                employee_id,
            )
            by_employee_day[key].append(row)
        else:
            unkeyed_rows.append(row)
    deduplicated_rows = list(unkeyed_rows)
    duplicate_observations = 0
    for group in by_employee_day.values():
        if len(group) > 1:
            duplicate_observations += len(group)
        else:
            deduplicated_rows.extend(group)
    if duplicate_observations:
        warnings.append(
            f"Excluded {duplicate_observations} attendance record(s) in duplicate employee/date groups because shift identity is unavailable."
        )

    eligible = 0
    present = 0
    excluded = 0
    unknown = 0
    for row in deduplicated_rows:
        status = _text(row.get("status")).casefold().replace("_", " ").replace("/", " ")
        compact_status = re.sub(r"[^a-z]", "", status)
        if (
            any(term in status for term in exclude_terms)
            or status in {"off", "weekly off", "holiday", "wo", "nm", "na", "not applicable"}
            or compact_status in {"al", "cl", "sl", "pl", "co"}
        ):
            excluded += 1
            continue
        if not status or any(term in status for term in ("half day", "half-day", "halfday", "partial")):
            unknown += 1
            continue
        is_present = any(term in status for term in present_terms) or compact_status in {"p", "pr", "l"}
        is_absent = (
            status in {"absent", "no show", "noshow", "a", "ab", "abs"}
            or compact_status in {"a", "ab", "abs"}
            or status.startswith("absent ")
        )
        if not is_present and not is_absent:
            unknown += 1
            continue
        eligible += 1
        if is_present:
            present += 1
    if excluded:
        warnings.append(f"Attendance denominator excludes {excluded} leave/day-off/not-marked record(s).")
    if unknown:
        warnings.append(f"Attendance denominator excludes {unknown} record(s) with unknown or partial-day status.")
    warnings.append("No scheduled-shift assumption is made; attendance rate uses only marked, non-excluded attendance records.")
    rate = present * 100.0 / eligible if eligible else None
    return rate, _sum_complete(deduplicated_rows, "hours")


def _metrics_for_rows(
    rows: list[dict[str, Any]],
    attendance_rows: list[dict[str, Any]],
    productivity_rows: list[dict[str, Any]],
    roster: list[dict[str, Any]],
    roster_known: bool,
    filters: dict[str, Any],
    selected_employee: dict[str, Any] | None,
    selected_role: str | None,
    warnings: list[str],
) -> dict[str, Any]:
    actual = _sum_complete(rows, "actual")
    budget = _sum_complete(rows, "budget")
    last_year = _sum_complete(rows, "last_year")
    if rows:
        if actual is None:
            warnings.append("Actual is unavailable because at least one selected sales row has no numeric actual value.")
        if budget is None:
            warnings.append("Budget is unavailable because at least one selected sales row has no numeric budget denominator.")
        if last_year is None:
            warnings.append("Last-year sales is unavailable because at least one selected sales row has no numeric value.")
    budget_pct = _percentage(actual, budget)
    growth_pct = None
    if actual is not None and last_year not in (None, 0):
        growth_pct = (actual - last_year) * 100.0 / last_year

    oos_value = _sum_complete(rows, "oos_value")
    if len(rows) == 1:
        oos_pct = _number(rows[0].get("oos_pct"))
    else:
        oos_pct = None
        if rows and any(_number(row.get("oos_pct")) is not None for row in rows):
            warnings.append("OOS percentage is unavailable for aggregated rows because a valid weighting denominator was not provided; OOS is not estimated as lost sales.")
    if any(
        _number(row.get("oos_value")) is not None or _number(row.get("oos_pct")) is not None
        for row in rows
    ):
        warnings.append("OOS values are source-reported; no lost-sales amount is estimated.")

    attendance_pct, hours = _attendance_metrics(attendance_rows, warnings) if attendance_rows else (None, None)
    if not attendance_rows:
        warnings.append("No attendance records are available for this period and scope.")
        warnings.append("No scheduled-shift assumption is made; no attendance denominator is inferred.")

    scoped_roster = [person for person in roster if _scope_matches(person, filters)]
    if roster_known:
        employee_count = len({
            (
                _text(person.get("employee_id") or person.get("employee_code")).casefold()
                or f"{_text(person.get('employee_name')).casefold()}|{_scope_value(person, 'store').casefold()}"
            )
            for person in scoped_roster
        })
    else:
        employee_count = None
        warnings.append("Employee count is unavailable because no as-of roster snapshot was available.")

    direct_reports = None
    if selected_employee and selected_role:
        manager_id = _text(selected_employee.get("employee_id") or selected_employee.get("employee_code"))
        manager_store = _scope_value(selected_employee, "store").casefold()
        direct_reports = len([
            person for person in roster
            if _text(person.get("manager_id")).casefold() == manager_id.casefold()
            and (
                not manager_store
                or _scope_value(person, "store").casefold() == manager_store
            )
        ]) if roster_known else None

    productivity = None
    if len(productivity_rows) == 1:
        productivity = _number(productivity_rows[0].get("productivity"))
    elif len(productivity_rows) > 1:
        warnings.append("Productivity is shown only when a single source-reported value is available; it is not averaged or estimated.")
    if productivity is None and productivity_rows:
        warnings.append("Productivity is unavailable because the source value is missing or non-numeric.")

    return {
        "actual": actual,
        "budget": budget,
        "last_year": last_year,
        "budget_achievement_pct": budget_pct,
        "sales_growth_pct": growth_pct,
        "oos_value": oos_value,
        "oos_pct": oos_pct,
        "attendance_rate_pct": attendance_pct,
        "hours": hours,
        "employee_count": employee_count,
        "direct_reports": direct_reports,
        "productivity": productivity,
    }


def _trend_rows(
    rows: list[dict[str, Any]],
    kind: str,
    warnings: list[str] | None = None,
) -> list[dict[str, Any]]:
    if kind == "sales_mtd":
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in rows:
            observation_period = _snapshot_period(row)
            parsed = _period_date(observation_period)
            period = parsed.strftime("%Y-%m") if parsed else observation_period[:7]
            if period:
                grouped[period].append(row)
    else:
        grouped = defaultdict(list)
        for row in rows:
            period = _text(row.get("date") or row.get("_period"))
            if period:
                grouped[period[:10]].append(row)
    result = []
    for period in sorted(grouped):
        day_rows = (
            _collapse_mtd_hierarchy(grouped[period], warnings)
            if kind == "sales_mtd"
            else _collapse_by_observation(grouped[period], kind)
        )
        actual = _sum_complete(day_rows, "actual")
        budget = _sum_complete(day_rows, "budget")
        last_year = _sum_complete(day_rows, "last_year")
        result.append({
            "period": period,
            "kind": kind,
            "actual": actual,
            "budget": budget,
            "last_year": last_year,
            "budget_achievement_pct": _percentage(actual, budget),
            "sales_growth_pct": (
                (actual - last_year) * 100.0 / last_year
                if actual is not None and last_year not in (None, 0)
                else None
            ),
        })
    return result


def analyze(store: Any, filters: dict[str, Any] | None = None) -> dict[str, Any]:
    """Analyze active local records without assuming missing values are zero."""
    requested = dict(filters or {})
    kind = _text(requested.get("kind")) or "sales_mtd"
    if kind not in {"sales_mtd", "sales_daily", "productivity", "attendance"}:
        kind = "sales_mtd"
    warnings: list[str] = []

    as_of = _text(requested.get("period_to")) or None
    if as_of and re.fullmatch(r"\d{4}-\d{2}", as_of):
        try:
            parsed_month = date.fromisoformat(f"{as_of}-01")
            as_of = date(parsed_month.year, parsed_month.month, calendar.monthrange(parsed_month.year, parsed_month.month)[1]).isoformat()
        except ValueError:
            pass
    base_roster, roster_known = _roster_snapshot(store, as_of)
    selected_employee = None
    selected_role = None
    employee_filter = _text(requested.get("employee"))
    if employee_filter:
        selected_employee, matches = _find_employee(base_roster, employee_filter)
        if len(matches) > 1:
            warnings.append(f"Employee '{employee_filter}' is ambiguous in the as-of roster; use a unique employee code.")
        elif selected_employee is None:
            warnings.append(f"Employee '{employee_filter}' was not found in the as-of roster.")
        else:
            selected_role = _manager_role(selected_employee)
            if selected_role is None:
                warnings.append("Employee sales are available only for observations when the as-of roster assigns GM, DH, or SM responsibility; ordinary-staff observations are excluded.")
            else:
                # Validate the current as-of assignment and explicit scope
                # overrides for transparent warnings. Historical records are
                # still independently resolved against their own roster dates.
                _apply_manager_scope(selected_employee, selected_role, requested, warnings)

    sales_rows: list[dict[str, Any]] = []
    if kind in {"sales_mtd", "sales_daily"}:
        sales_rows = _get_rows(store, kind, requested)
        sales_rows = [row for row in sales_rows if _scope_matches(row, requested)]
        if kind == "sales_mtd":
            sales_rows = _latest_mtd_snapshots(sales_rows)
        if employee_filter and selected_employee is None:
            sales_rows = []
        elif employee_filter:
            sales_rows = _filter_employee_observations(
                store, sales_rows, selected_employee, kind, requested, warnings
            )
        if kind == "sales_mtd":
            sales_rows = _collapse_mtd_hierarchy(sales_rows, warnings)
        else:
            sales_rows = _collapse_by_observation(sales_rows, kind)

    attendance_rows = _get_rows(store, "attendance", requested)
    attendance_rows = [
        row for row in attendance_rows
        if _scope_matches(row, requested)
    ]
    if employee_filter and selected_employee is None:
        attendance_rows = []
    elif employee_filter:
        attendance_rows = _filter_employee_observations(
            store,
            attendance_rows,
            selected_employee,
            "attendance",
            requested,
            warnings,
            allow_ordinary_employee=True,
        )
    attendance_rows = _collapse_by_observation(attendance_rows, "attendance")

    productivity_rows = _get_rows(store, "productivity", requested)
    productivity_rows = [row for row in productivity_rows if _scope_matches(row, requested)]
    if kind == "sales_mtd":
        productivity_rows = _latest_mtd_snapshots(productivity_rows)
    if employee_filter and selected_employee is None:
        productivity_rows = []
    elif employee_filter:
        productivity_rows = _filter_employee_observations(
            store,
            productivity_rows,
            selected_employee,
            "sales_mtd" if kind == "sales_mtd" else "productivity",
            requested,
            warnings,
            allow_ordinary_employee=True,
        )
    if kind == "sales_mtd":
        productivity_rows = _collapse_mtd_hierarchy(productivity_rows, warnings)
    else:
        productivity_rows = _collapse_by_observation(productivity_rows, "productivity")

    roster_for_scope = base_roster
    roster_scope_known = roster_known
    if selected_employee and selected_role:
        assignment = {key: _scope_value(selected_employee, key) for key in ("store", "department", "section")}
        required = ("store",) if selected_role == "GM" else ("store", "department") if selected_role == "DH" else ("store", "department", "section")
        if any(not assignment[key] for key in required):
            roster_for_scope = []
            roster_scope_known = False
        else:
            scope_values = dict(requested)
            for key in required:
                scope_values[key] = assignment[key]
            roster_for_scope = [person for person in base_roster if _scope_matches(person, scope_values)]
    elif selected_employee:
        roster_for_scope = [selected_employee]

    metrics = _metrics_for_rows(
        sales_rows,
        attendance_rows,
        productivity_rows,
        roster_for_scope,
        roster_scope_known,
        requested,
        selected_employee,
        selected_role,
        warnings,
    )
    if not sales_rows and kind in {"sales_mtd", "sales_daily"}:
        warnings.append(f"No active {kind} records match the selected period and scope.")
    if selected_employee:
        warnings.append("Employee sales scope is resolved from the roster as of each observation using GM/DH/SM responsibility only.")
    if kind == "sales_mtd":
        warnings.append("MTD analysis uses the latest available snapshot per month and scope; cumulative snapshots are not summed.")
    elif kind == "sales_daily":
        warnings.append("Daily analysis uses daily records only; MTD snapshots are not mixed into daily totals.")

    commentary: list[str] = []
    for key, label, suffix in (
        ("actual", "Actual", ""),
        ("budget_achievement_pct", "Budget achievement", "%"),
        ("sales_growth_pct", "Sales growth", "%"),
        ("attendance_rate_pct", "Attendance rate", "%"),
    ):
        value = metrics[key]
        commentary.append(f"{label}: unavailable." if value is None else f"{label}: {value:,.2f}{suffix}.")
    commentary.append("Productivity is the source-reported value; it is not independently recalculated.")
    commentary.append("No authoritative performance score or automated employment decision is produced.")

    evidence = {
        "actual": {"value": metrics["actual"], "evidence": "Sum of selected non-overlapping source actual values."},
        "budget": {"value": metrics["budget"], "evidence": "Sum of selected non-overlapping source budget values; incomplete source values remain unavailable."},
        "last_year": {"value": metrics["last_year"], "evidence": "Sum of selected non-overlapping source last-year values; incomplete source values remain unavailable."},
        "budget_achievement_pct": {
            "value": metrics["budget_achievement_pct"],
            "evidence": "100 × aggregated actual ÷ aggregated budget; unavailable without a known non-zero budget denominator.",
        },
        "sales_growth_pct": {
            "value": metrics["sales_growth_pct"],
            "evidence": "100 × (actual − last year) ÷ last year; unavailable without a known non-zero last-year denominator.",
        },
        "attendance_rate_pct": {
            "value": metrics["attendance_rate_pct"],
            "evidence": "Marked present/late observations ÷ marked non-excluded present/late/absent observations; partial or unknown statuses are excluded.",
        },
        "hours": {"value": metrics["hours"], "evidence": "Sum of complete source-reported hours from the selected attendance records."},
        "employee_count": {"value": metrics["employee_count"], "evidence": "Unique employees in the latest as-of roster snapshot within scope."},
        "direct_reports": {"value": metrics["direct_reports"], "evidence": "As-of roster records whose manager_id equals the selected manager's employee_id, within the same store."},
        "oos_value": {"value": metrics["oos_value"], "evidence": "Source-reported OOS value; never treated as estimated lost sales."},
        "oos_pct": {"value": metrics["oos_pct"], "evidence": "Source percentage for a single selected row only; aggregation requires an explicit known denominator."},
        "productivity": {
            "value": metrics["productivity"],
            "evidence": "Source-reported value only; no aggregation or estimation.",
        },
    }
    score = {
        "status": "PROVISIONAL",
        "score": None,
        "authoritative": False,
        "approved_criteria": False,
        "components": evidence,
        "note": "No approved full scoring criteria were supplied; components are evidence only and must not drive automated employment decisions.",
    }
    relevant_rows = sales_rows + attendance_rows + productivity_rows + roster_for_scope
    return {
        "metrics": metrics,
        "commentary": commentary,
        "warnings": list(dict.fromkeys(warnings)),
        "trends": _trend_rows(sales_rows, kind, warnings) if kind in {"sales_mtd", "sales_daily"} else [],
        "rows": sales_rows,
        "sources": _source_records(relevant_rows),
        "score": score,
        "filters": {
            key: requested.get(key)
            for key in ("kind", "period_from", "period_to", "country", "store", "department", "section", "employee")
            if requested.get(key) is not None
        },
    }


def _report_values(result: dict[str, Any], filters: dict[str, Any] | None) -> list[tuple[str, str]]:
    selected_filters = filters if filters is not None else result.get("filters", {})
    values: list[tuple[str, str]] = []
    values.append(("Report", "PNO offline analysis"))
    values.append(("Filters", json.dumps(selected_filters or {}, ensure_ascii=False, sort_keys=True, default=str)))
    for key in METRIC_NAMES:
        values.append((key, "" if result.get("metrics", {}).get(key) is None else str(result["metrics"][key])))
    for index, item in enumerate(result.get("commentary", []), 1):
        values.append((f"Commentary {index}", str(item)))
    for index, item in enumerate(result.get("warnings", []), 1):
        values.append((f"Warning {index}", str(item)))
    for index, item in enumerate(result.get("sources", []), 1):
        values.append((f"Source {index}", json.dumps(item, ensure_ascii=False, sort_keys=True, default=str)))
    return values


def _csv_safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    if re.match(r"^[\s]*[=+\-@]", value):
        return "'" + value
    return value


def _append_xlsx_row(sheet: Any, values: Iterable[Any]) -> None:
    row_index = sheet.max_row
    if row_index == 1 and all(cell.value is None for cell in sheet[1]):
        row_index = 1
    else:
        row_index += 1
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row=row_index, column=column)
        cell.value = value
        if isinstance(value, str):
            # openpyxl otherwise treats leading '=' as a formula cell.
            cell.data_type = "s"


def export_report(
    result: dict[str, Any],
    path: str | Path,
    format: str = "pdf",
    filters: dict[str, Any] | None = None,
) -> Path:
    """Export a report including its filters, source provenance, and warnings."""
    target = Path(path)
    chosen = _text(format).lower().lstrip(".") or target.suffix.lower().lstrip(".")
    if chosen not in {"pdf", "xlsx", "csv"}:
        raise ValueError("Unsupported report format; choose pdf, xlsx, or csv.")
    target.parent.mkdir(parents=True, exist_ok=True)
    values = _report_values(result, filters)

    if chosen == "csv":
        with target.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(("item", "value"))
            writer.writerows(tuple(_csv_safe(value) for value in row) for row in values)
            writer.writerow(())
            writer.writerow(tuple(_csv_safe(value) for value in ("trend_period", "kind", "actual", "budget", "last_year", "budget_achievement_pct", "sales_growth_pct")))
            for trend in result.get("trends", []):
                writer.writerow(tuple(
                    _csv_safe(trend.get(key))
                    for key in ("period", "kind", "actual", "budget", "last_year", "budget_achievement_pct", "sales_growth_pct")
                ))
        return target

    if chosen == "xlsx":
        try:
            from openpyxl import Workbook
        except ImportError as error:
            raise RuntimeError("XLSX export requires the optional openpyxl package.") from error
        workbook = Workbook()
        overview = workbook.active
        overview.title = "Overview"
        _append_xlsx_row(overview, ("Item", "Value"))
        for item in values:
            _append_xlsx_row(overview, item)
        trends = workbook.create_sheet("Trends")
        trend_keys = ("period", "kind", "actual", "budget", "last_year", "budget_achievement_pct", "sales_growth_pct")
        _append_xlsx_row(trends, trend_keys)
        for trend in result.get("trends", []):
            _append_xlsx_row(trends, [trend.get(key) for key in trend_keys])
        sources = workbook.create_sheet("Sources")
        source_keys = ("import_id", "kind", "period", "source", "sheet")
        _append_xlsx_row(sources, source_keys)
        for source in result.get("sources", []):
            _append_xlsx_row(sources, [source.get(key) for key in source_keys])
        workbook.save(target)
        return target

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError as error:
        raise RuntimeError("PDF export requires the optional reportlab package.") from error
    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph("PNO offline analysis", styles["Title"]), Spacer(1, 10)]
    table_rows = [["Item", "Value"]]
    table_rows.extend([[label, value] for label, value in values])
    table = Table(table_rows, colWidths=(135, 390), repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#31261D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F2EA")]),
    ]))
    story.append(table)
    if result.get("trends"):
        story.extend([Spacer(1, 12), Paragraph("Historical trend", styles["Heading2"])])
        trend_rows = [["Period", "Kind", "Actual", "Budget", "Last year", "Achievement %", "Growth %"]]
        for trend in result["trends"]:
            trend_rows.append([str(trend.get(key) if trend.get(key) is not None else "") for key in (
                "period", "kind", "actual", "budget", "last_year", "budget_achievement_pct", "sales_growth_pct"
            )])
        trend_table = Table(trend_rows, repeatRows=1)
        trend_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#B5975A")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]))
        story.append(trend_table)
    SimpleDocTemplate(str(target), pagesize=letter).build(story)
    return target