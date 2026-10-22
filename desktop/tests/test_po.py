"""P&O workbook pack: import, insights, analyst and exports (fictional data only)."""

from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
from openpyxl import Workbook

from pno import canon, privacy
from pno.insights import Insights
from pno.po_assistant import PoAnalyst
from pno.storage import Store
from pno.workbook_pack import PACK_KINDS, import_pack, num, parse_sheet, scan_pack

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    spec = importlib.util.spec_from_file_location("make_samples", ROOT / "scripts" / "make_samples.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.make_po_workbook(tmp_path_factory.mktemp("pack") / "fictional-po.xlsx")


@pytest.fixture()
def loaded(workbook, tmp_path):
    store = Store(tmp_path / "pno.sqlite3")
    messages = import_pack(store, scan_pack(workbook), "2026-10")
    assert all(m.startswith("✓") for m in messages), messages
    return store


def test_every_tab_is_recognised(workbook):
    plan = scan_pack(workbook)
    assert not plan.skipped
    assert {s.kind for s in plan.sheets} == set(PACK_KINDS)
    master = next(s for s in plan.sheets if s.kind == "po_master")
    for forbidden in ("birth", "mobile", "email", "marital"):
        assert not any(forbidden in key.casefold() for key in master.records[0])
    assert "age" not in master.records[0]


def test_name_dictionary():
    assert canon.section("S050-S051 Deli & Dairy") == "Deli"
    assert canon.section("L.H.H.") == "Light Household"
    assert canon.section("Comm.FMCG") == "Consumer Goods"
    assert canon.section("Fruits/Veg") == "Fruits & Vegetables"
    assert canon.store("652 LAH High Street Paragon Ci") == "LAH Paragon City"
    assert canon.store("STC-Gujranwala") == "GUJ Steel Casting"
    assert canon.store("HB PK KCH H&B PAK KCH LUK (MYLI)") == "KCH Lucky One · Myli"
    assert canon.repair_name("MuhMORmad SMORi HMORza") == "Muhammad Sami Hamza"
    assert canon.sections_in("Leading Bakery & Delicatessen Sections") == ["Bakery", "Deli"]
    assert canon.shift("MOr") == "Morning" and canon.shift("DAY OFF") == "Off" and canon.shift("C/L") == "Leave"


def test_number_parsing():
    assert num("(4.1%)") == -4.1
    assert num("142,018,103") == 142018103
    assert num(" -   ") is None
    assert num("▼ -6%") == -6


def test_percent_formatted_cells_become_points(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Sections", "HM PK LAH Fortress"])
    ws.append(["", "Shrinkges", "", "", "Stock Days"])
    ws.append(["", "Budget %", "Actual %", "Act-Bud", "Budget", "Actual"])
    ws.append(["FFD - Bakery", -0.001, -0.004, -0.003, 15, 20])
    for cell in ws[4][1:4]:
        cell.number_format = "0.0%"
    path = tmp_path / "shrink.xlsx"
    wb.save(path)
    from pno.workbook_pack import read_sheets

    name, rows = read_sheets(path)[0]
    sheet = parse_sheet(name, rows, "po_shrink")
    record = sheet.records[0]
    assert record["section"] == "Bakery"
    assert record["shrink_actual_pct"] == pytest.approx(-0.4)
    assert record["stock_days_actual"] == 20


def test_reimporting_a_month_replaces_it(workbook, loaded):
    import_pack(loaded, scan_pack(workbook), "2026-10")
    active = [i for i in loaded.list_imports() if i["active"] and i["kind"] == "po_rating"]
    assert len(active) == 1


def test_insights_cover_all_views(loaded):
    insights = Insights(loaded)
    headline = insights.headline()
    assert headline["stores"] == 3
    assert headline["sales"] > 0
    cards = insights.store_scorecard()
    assert {c["status"] for c in cards} <= {"On track", "Watch", "At risk", "No data"}
    leaders = insights.leaders()
    assert any(l.score is not None for l in leaders)
    scored = [l.score for l in leaders if l.score is not None]
    assert all(0 <= s <= 100 for s in scored)
    assert scored == sorted(scored, reverse=True)
    assert insights.turnover_summary()["total"] > 0
    assert insights.roster_summary()["people"] == 15
    stores, sections, values = insights.section_matrix("sales_vs_budget")
    assert sections and any(v is not None for v in values.values())
    assert isinstance(insights.alerts(), list)


def test_missing_measures_are_not_zero(loaded):
    insights = Insights(loaded, weights={"rating": 0})
    for leader in insights.leaders():
        assert "rating" not in leader.components or insights.weights["rating"] == 0
        if len([k for k in leader.components if insights.weights.get(k)]) < 3:
            assert leader.score is None


def test_analyst_answers_from_data(loaded):
    analyst = PoAnalyst(Insights(loaded))
    assert "Country brief" in analyst.answer("Brief me for the country review")
    assert "Stores by productivity" in analyst.answer("rank stores by productivity")
    assert "leaders" in analyst.answer("who are the top 3 leaders?").lower()
    assert "Fortress" in analyst.answer("how is Fortress doing?")
    assert "Example Leader" in analyst.answer("Tell me about Example Leader 01")
    assert analyst.answer("hello there") is None


def test_presentation_mode_masks_names(loaded):
    from pno.po_format import leader_name

    insights = Insights(loaded)
    leader = insights.leaders()[0]
    privacy.set_demo_mode(True)
    try:
        masked = leader_name(leader)
        assert masked.startswith("Leader ") and "Example" not in masked
        assert "Example Leader" not in PoAnalyst(insights).answer("who are the top 5 leaders")
    finally:
        privacy.set_demo_mode(False)


def test_exports(loaded, tmp_path):
    from pno.po_reports import export_board_pack, export_leader_pdf, export_workbook

    insights = Insights(loaded)
    pdf = export_board_pack(insights, str(tmp_path / "pack.pdf"))
    assert Path(pdf).read_bytes()[:4] == b"%PDF"
    xlsx = export_workbook(insights, str(tmp_path / "data.xlsx"))
    from openpyxl import load_workbook

    assert "Store scorecard" in load_workbook(xlsx).sheetnames
    one = export_leader_pdf(insights, insights.leaders()[0], str(tmp_path / "one.pdf"))
    assert Path(one).stat().st_size > 1000


def test_old_filters_only_see_the_staff_roster(loaded):
    roster_stores = {r["store"] for r in loaded.records(kind="roster")}
    assert set(loaded.dimensions()["stores"]) == roster_stores
